#!/usr/bin/env python3
"""Smoke test des agents iziLife contre staging ou production, sans lire de .env.

Le runner compile les scripts, valide tous les ``--init``, audite les vraies
files du Drive ciblé, exécute éventuellement les agents, puis produit des
rapports JSON et Markdown PASS/FAIL.
"""
from __future__ import annotations

import argparse
import compileall
import json
import os
import re
import subprocess
import sys
import tempfile
import time
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl est requis: pip install openpyxl") from exc


IZILIFE_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = IZILIFE_ROOT.parents[1]
REPORT_DIR = Path(__file__).resolve().parent / "reports"


@dataclass(frozen=True)
class Agent:
    name: str
    script: str
    folder: str
    workbook: str


AGENTS = (
    Agent("territorial", "geography/classificateur_territoires.py", "geography", "territories.xlsx"),
    Agent("city_subdivisions", "geography/city_subdivisions.py", "geography", "city_subdivisions.xlsx"),
    Agent("event_structures", "events/event_structures.py", "events", "event_structures.xlsx"),
    Agent("event_calendar", "events/event_calendar.py", "events", "event_calendar.xlsx"),
    Agent("streets", "geography/streets.py", "geography", "streets.xlsx"),
    Agent("equipment", "equipments/equipment.py", "equipments", "equipments.xlsx"),
    Agent("escapade", "escapades/escapade_proposal.py", "escapades", "escapade_proposals.xlsx"),
    Agent("link", "links/link.py", "links", "links.xlsx"),
    Agent("spotlight", "spotlights/spotlight.py", "spotlights", "spotlights.xlsx"),
    Agent("tags", "tags/tags.py", "tags", "tags.xlsx"),
    Agent("activity", "activity_proposals/activity_proposal.py", "activity-proposals", "activity_proposals.xlsx"),
    Agent("brunch", "brunch/brunch.py", "brunch", "brunches.xlsx"),
    Agent("special-offer", "offers/special_offer.py", "offers", "special_offers.xlsx"),
    Agent("subscription-benefit", "offers/izilife_subscription_benefit.py", "offers", "izilife_subscription_benefits.xlsx"),
    Agent("enrichisseur++", "places/enrichisseur_contenus.py", "places", "enrichisseur_contenus.xlsx"),
)


@dataclass
class Result:
    name: str
    phase: str
    state: str
    seconds: float = 0.0
    detail: str = ""
    command: list[str] | None = None


def normalized_zone(value: str) -> str:
    value = value.strip().lower()
    return value if value.endswith("-zone") else f"{value}-zone"


def drive_root(env_name: str) -> Path:
    base = os.getenv("AGENTIC_DRIVE_ROOT")
    if not base:
        raise RuntimeError("AGENTIC_DRIVE_ROOT doit être fourni par le processus")
    suffix = {
        "local": "agentic_workspace_local",
        "staging": "agentic_workspace_staging",
        "prod": "agentic_workspace",
    }[env_name]
    return Path(base).expanduser() / suffix / "izilife"


def workbook_path(agent: Agent, env_name: str, zone: str) -> Path:
    return drive_root(env_name) / agent.folder / normalized_zone(zone) / agent.workbook


def command(agent: Agent, zone: str, env_name: str, *extra: str) -> list[str]:
    return [sys.executable, str(IZILIFE_ROOT / agent.script), "--zone", zone, "--env", env_name, *extra]


def run_command(name: str, phase: str, cmd: list[str], env: dict[str, str] | None = None) -> Result:
    started = time.monotonic()
    proc = subprocess.run(cmd, cwd=IZILIFE_ROOT, env=env, text=True, capture_output=True)
    elapsed = time.monotonic() - started
    output = "\n".join(part.strip() for part in (proc.stdout, proc.stderr) if part.strip())
    return Result(name, phase, "PASS" if proc.returncode == 0 else "FAIL", elapsed, output[-6000:], cmd)


def prepare_fixture(agent: Agent, test_zone: str, env_name: str) -> Result:
    target = workbook_path(agent, env_name, test_zone)
    if not target.exists():
        return Result(agent.name, "prepare", "SKIP", detail=f"Classeur {env_name} absent: {target}")
    return Result(agent.name, "prepare", "PASS", detail=f"Classeur {env_name} existant conservé: {target}")


def pending_rows(path: Path, sheets: list[str] | None = None) -> tuple[int, dict[str, int]]:
    if not path.exists(): return 0, {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    counts: dict[str, int] = {}
    for ws in wb.worksheets:
        if sheets and ws.title not in sheets: continue
        first = next(ws.iter_rows(min_row=1, max_row=1), ())
        headers = [str(c.value or "").strip().lower() for c in first]
        status_name = "status" if "status" in headers else "statut" if "statut" in headers else ""
        if not status_name: continue
        col = headers.index(status_name)
        counts[ws.title] = sum(
            1
            for row in ws.iter_rows(min_row=2, values_only=True)
            if any(str(value or "").strip() for value in row)
            and str(row[col] or "").strip().lower() in {"", "pending", "retry", "error", "failed"}
        )
    return sum(counts.values()), counts


def audit_inputs(zone: str, env_name: str) -> list[Result]:
    root = drive_root(env_name)
    z = normalized_zone(zone)
    requirements = (
        ("manual-events", root / "events" / z / "curate_events.xlsx", ["Shotgun", "Facebook", "HelloAsso", "Eventbrite", "Meetup", "Instagram"], "1 à 5 liens manuels Shotgun, Facebook, HelloAsso, Eventbrite, Meetup ou Instagram"),
        ("event-structures", root / "events" / z / "event_structures.xlsx", ["Structures"], "une EventSerie, une AnnualCelebration et un événement récurrent avec leurs informations éditoriales et calendaires"),
        ("event-calendar", root / "events" / z / "event_calendar.xlsx", ["Calendar"], "une opération datée maîtrisée sur EventSerie, AnnualCelebration ou Event"),
        ("media-downloads", root / "medias" / "medias_igram.xlsx", ["Medias"], "liens de médias à télécharger"),
        ("territorial", root / "geography" / z / "territories.xlsx", ["Cities", "AdministrativeDivisions"], "villes/divisions à configurer"),
        ("city-subdivisions", root / "geography" / z / "city_subdivisions.xlsx", ["Subdivisions"], "un quartier à créer dans une ville staging existante"),
        ("streets", root / "geography" / z / "streets.xlsx", ["Streets"], "une rue connue et au moins un membre SHOP ou PLACE existant"),
        ("enrichisseur", root / "places" / z / "ameliorateur.xlsx", ["Prioritaires"], "nouveaux string_id de lieux à améliorer"),
        ("enrichisseur++", root / "places" / z / "enrichisseur_contenus.xlsx", ["Sources"], "sources produits, services, tarifs/EAP et équipements"),
        ("equipment", root / "equipments" / z / "equipments.xlsx", ["Equipments"], "un équipement avec catégorie et caractéristiques spécialisées"),
        ("escapade", root / "escapades" / z / "escapade_proposals.xlsx", ["EscapadeProposals"], "une proposition éditoriale sur un scope existant"),
        ("links", root / "links" / z / "links.xlsx", ["Links"], "un lien sur un scope existant"),
        ("spotlight", root / "spotlights" / z / "spotlights.xlsx", ["Spotlights"], "un Spotlight éditorial, avec GOOD_FOR uniquement si justifié"),
        ("tags", root / "tags" / z / "tags.xlsx", ["Tags"], "une affectation et/ou un retrait de tag sur un scope existant"),
        ("activity-proposals", root / "activity-proposals" / z / "activity_proposals.xlsx", ["ActivityProposals"], "une proposition d'activité sur un lieu existant"),
        ("brunch", root / "brunch" / z / "brunches.xlsx", ["Brunchs"], "une configuration de brunch sur un SHOP ou PLACE existant"),
        ("special-offer", root / "offers" / z / "special_offers.xlsx", ["Offres"], "une Offer avec sa BPR et son Effect"),
        ("subscription-benefit", root / "offers" / z / "izilife_subscription_benefits.xlsx", ["Avantages abonnements"], "une récompense iziLife : BPR de plan avec son Effect"),
    )
    results = []
    for name, path, sheets, request in requirements:
        count, detail = pending_rows(path, sheets)
        state = "PASS" if count else "NEEDS_INPUT"
        message = f"{count} ligne(s) prêtes {detail} | {path}" if count else f"À fournir : {request} | {path}"
        results.append(Result(name, "inputs", state, detail=message))
    return results


def inspect_workbook(agent: Agent, zone: str, env_name: str) -> Result:
    path = workbook_path(agent, env_name, zone)
    if not path.exists():
        return Result(agent.name, "verify", "SKIP", detail=f"Fixture absente: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    totals: dict[str, dict[str, int]] = {}
    bad = 0
    for ws in wb.worksheets:
        headers = [str(c.value or "").strip().lower() for c in ws[1]]
        if "status" not in headers:
            continue
        col = headers.index("status") + 1
        states: dict[str, int] = {}
        for row in range(2, ws.max_row + 1):
            state = str(ws.cell(row, col).value or "pending").strip().lower()
            states[state] = states.get(state, 0) + 1
            bad += int(state in {"error", "unmapped"})
        totals[ws.title] = states
    return Result(agent.name, "verify", "FAIL" if bad else "PASS", detail=json.dumps(totals, ensure_ascii=False))


def init_checks(zone: str, agents: list[Agent]) -> list[Result]:
    results: list[Result] = []
    with tempfile.TemporaryDirectory(prefix="izilife-agent-init-") as tmp:
        env = os.environ.copy()
        env["AGENTIC_DRIVE_ROOT"] = tmp
        for agent in agents:
            cmd = [sys.executable, str(IZILIFE_ROOT / agent.script), "--zone", zone, "--env", "local", "--init"]
            result = run_command(agent.name, "init", cmd, env)
            expected = Path(tmp) / "agentic_workspace_local" / "izilife" / agent.folder / normalized_zone(zone) / agent.workbook
            if result.state == "PASS" and not expected.exists():
                result.state = "FAIL"
                result.detail += f"\nClasseur attendu non créé: {expected}"
            results.append(result)
    return results


def write_report(results: list[Result], args) -> tuple[Path, Path]:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    json_path = REPORT_DIR / f"{args.target_env}-smoke-{stamp}.json"
    md_path = REPORT_DIR / f"{args.target_env}-smoke-{stamp}.md"
    payload = {"generated_at": datetime.now().isoformat(), "arguments": vars(args), "results": [asdict(r) for r in results]}
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = [f"# iziLife {args.target_env} smoke test", "", f"Généré : {payload['generated_at']}", "", "| Phase | Agent | État | Durée | Détail |", "|---|---|---:|---:|---|"]
    for result in results:
        detail = result.detail.replace("\n", "<br>").replace("|", "\\|")
        lines.append(f"| {result.phase} | {result.name} | {result.state} | {result.seconds:.1f}s | {detail} |")
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return json_path, md_path


def parse_args():
    parser = argparse.ArgumentParser(description="Smoke tests reproductibles des agents iziLife.")
    parser.add_argument("--test-zone", default="lille", help="Zone normale utilisée par les agents.")
    parser.add_argument("--target-env", choices=["staging", "prod"], default="staging")
    parser.add_argument("--allow-prod", action="store_true", help="Confirmation obligatoire pour --run --target-env prod.")
    parser.add_argument("--max", type=int, default=3, choices=range(1, 6), metavar="1..5")
    parser.add_argument("--prepare", action="store_true", help="Auditer les files Excel du Drive et indiquer les données à fournir.")
    parser.add_argument("--run", action="store_true", help="Autoriser les écritures via les API de l'environnement ciblé.")
    parser.add_argument("--online", action="store_true", help="Tester les deux améliorateurs (max N) et la collecte Chasseur (3 minutes, N/catégorie).")
    parser.add_argument("--event-search", action="store_true", help="Tester la recherche automatique d'événements, sauf Facebook, avec --max-events.")
    parser.add_argument("--media", action="store_true", help="Tester la file de téléchargement des médias avec --max.")
    parser.add_argument("--only", nargs="*", choices=[a.name for a in AGENTS], help="Limiter la liste d'agents.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.run and args.target_env == "prod" and not args.allow_prod:
        raise SystemExit("Refus d'écrire en prod sans --allow-prod")
    chosen = [a for a in AGENTS if not args.only or a.name in args.only]
    results: list[Result] = []

    compiled = compileall.compile_dir(IZILIFE_ROOT, quiet=1, rx=re.compile(r"[\\/]old[\\/]"))
    results.append(Result("python", "compile", "PASS" if compiled else "FAIL", detail=str(IZILIFE_ROOT)))
    results.extend(init_checks(args.test_zone, chosen))

    if args.prepare:
        try:
            results.extend(prepare_fixture(a, args.test_zone, args.target_env) for a in chosen)
            results.extend(audit_inputs(args.test_zone, args.target_env))
        except Exception as exc:
            results.append(Result("fixtures", "prepare", "FAIL", detail=str(exc)))

    if args.run:
        if not (os.getenv("IZILIFE_AGENT_TOKEN") or os.getenv("AGENT_SECRET_TOKEN")):
            results.append(Result("credentials", "run", "FAIL", detail="Jeton absent du processus; aucun .env n'a été lu."))
        else:
            for agent in chosen:
                if args.online and agent.name == "enrichisseur++":
                    continue
                try:
                    path = workbook_path(agent, args.target_env, args.test_zone)
                except Exception as exc:
                    results.append(Result(agent.name, "run", "FAIL", detail=str(exc)))
                    continue
                if not path.exists():
                    results.append(Result(agent.name, "run", "SKIP", detail=f"Fixture absente: {path}"))
                    continue
                results.append(run_command(agent.name, "run", command(agent, args.test_zone, args.target_env, "--max", str(args.max))))
                results.append(inspect_workbook(agent, args.test_zone, args.target_env))

            if args.online:
                city = "lille"
                enrich_cmd = [sys.executable, str(IZILIFE_ROOT / "places/enrichisseur.py"), "--zone", args.test_zone, "--city", city, "--env", args.target_env, "--priority-only", "--max", str(args.max)]
                results.append(run_command("enrichisseur", "online", enrich_cmd))
                content_cmd = [sys.executable, str(IZILIFE_ROOT / "places/enrichisseur_contenus.py"), "--zone", args.test_zone, "--env", args.target_env, "--max", str(args.max)]
                results.append(run_command("enrichisseur++", "online", content_cmd))
                hunter_cmd = [sys.executable, str(IZILIFE_ROOT / "places/chasseur_lieux.py"), "--zone", args.test_zone, "--city", city, "--env", args.target_env, "--collect-only", "--max-per-cat", str(args.max), "--max-duration", "3"]
                results.append(run_command("chasseur", "online", hunter_cmd))

            if args.event_search:
                event_scripts = (
                    ("shotgun", "events/shotgun_scraper.py", False),
                    ("meetup", "events/meetup_scraper.py", True),
                    ("helloasso", "events/helloasso_scraper.py", True),
                    ("billetweb", "events/billetweb_scraper.py", True),
                    ("eventbrite", "events/eventbrite_scraper.py", True),
                )
                for name, script, with_zone in event_scripts:
                    cmd = [sys.executable, str(IZILIFE_ROOT / script), "--env", args.target_env, "--city", "lille", "--pages", "1", "--max-events", str(args.max)]
                    if with_zone: cmd.extend(["--zone", "lille"])
                    results.append(run_command(name, "events", cmd))

            if args.media:
                media_cmd = [sys.executable, str(IZILIFE_ROOT / "media_downloader.py"), "--env", args.target_env, "--max", str(args.max)]
                results.append(run_command("media-downloads", "media", media_cmd))

    json_path, md_path = write_report(results, args)
    for result in results:
        print(f"[{result.state}] {result.phase}/{result.name} {result.detail[:180]}")
    print(f"Rapports: {md_path} | {json_path}")
    return 1 if any(r.state == "FAIL" for r in results) else 0


if __name__ == "__main__":
    raise SystemExit(main())
