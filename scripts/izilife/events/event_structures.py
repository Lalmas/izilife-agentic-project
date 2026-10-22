#!/usr/bin/env python3
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "offers"))
from offer_agent_common import create_workbook, drive_workspace_root, normalize_zone, process

FILE, SHEET = "event_structures.xlsx", "Structures"
COLUMNS = [
    ("object_type", "scope"), ("string_id", "scope"), ("unique_id", "scope"),
    ("name", "offer"), ("description", "offer"), ("event_category_ref", "offer"),
    ("is_active", "bpr"), ("listable", "bpr"),
    ("partner_ref", "scope"), ("page_ref", "scope"),
    ("location_ref", "scope"), ("default_location_ref", "scope"),
    ("default_address", "scope"), ("default_address_complement", "scope"),
    ("default_zip_code", "scope"), ("default_city", "scope"),
    ("default_city_ref", "scope"), ("default_latitude", "scope"), ("default_longitude", "scope"),
    ("scope_ref", "scope"), ("celebration_day_ref", "scope"),
    ("can_be_linked_by_others", "bpr"), ("is_internal", "bpr"),
    ("news_visibility_days_before", "bpr"),
    ("is_izilife_serie", "bpr"), ("is_exclusive_on_izilife", "bpr"),
    ("is_izilife_partnership_serie", "bpr"), ("is_izilife_organizer_serie", "bpr"),
    ("access_type_ref", "access"), ("booking_needed_ref", "access"),
    ("access_on_booking", "access"), ("access_text", "access"), ("access_link", "access"),
    ("basic_price", "access"), ("currency_ref", "access"),
    ("principal_activity_category_ref", "offer"), ("default_event_category_ref", "offer"),
    ("first_title", "offer"), ("first_start_date", "access"), ("first_start_hour", "access"),
    ("first_end_date", "access"), ("first_end_hour", "access"),
    ("start_date", "access"), ("start_hour", "access"),
    ("end_date", "access"), ("end_hour", "access"),
    ("recurrence_frequency", "access"), ("recurrence_interval", "access"),
    ("recurrence_days", "access"), ("recurrence_month_day", "access"),
    ("recurrence_until", "access"), ("recurrence_rrule", "access"),
    ("timezone_id", "access"), ("horizon_days", "access"), ("count_limit", "access"),
    ("status", "result"), ("notes", "result"), ("last_result", "result"),
]


def payload(row: dict) -> dict:
    data = {
        key: row.get(key)
        for key, _ in COLUMNS
        if key not in {"status", "notes", "last_result"}
        and row.get(key) not in (None, "")
    }
    unique_id = str(data.get("unique_id") or "").strip()
    if str(data.get("object_type") or "").strip().upper() == "RECURRENT_EVENT" and len(unique_id) > 32:
        raise ValueError("unique_id RECURRENT_EVENT dépasse 32 caractères; la base le tronquerait.")
    return data


def add_guide(path: Path) -> None:
    wb = openpyxl.load_workbook(path)
    if "Guide" in wb.sheetnames:
        return
    ws = wb.create_sheet("Guide")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 115
    rows = [
        ("Principe", "Une ligne crée ou met à jour une EventSerie, une AnnualCelebration ou le parent d'un événement récurrent. Relancer la même ligne ne crée pas de doublon."),
        ("object_type", "EVENT_SERIE | ANNUAL_CELEBRATION | RECURRENT_EVENT"),
        ("Identité", "EVENT_SERIE et ANNUAL_CELEBRATION : string_id. RECURRENT_EVENT : unique_id après la première création ; string_id peut servir à l'upsert initial."),
        ("Références", "Utiliser les string_id ou l'identifiant numérique : event_category_ref, access_type_ref, currency_ref, etc."),
        ("Lieux", "location_ref pour RECURRENT_EVENT ; default_location_ref pour les deux structures. Syntaxe SHOP:soultraincafe ou PLACE:mon-lieu."),
        ("Portée annuelle", "scope_ref accepte CITY:lille, COUNTRY:france ou AREA:..."),
        ("Première date", "Pour EVENT_SERIE / ANNUAL_CELEBRATION, first_start_* et first_end_* créent facultativement la première édition."),
        ("Récurrence", "Pour RECURRENT_EVENT, renseigner start/end puis frequency DAILY/WEEKLY/MONTHLY/YEARLY, interval et éventuellement recurrence_days (MO,TU,WE,TH,FR,SA,SU)."),
        ("RRULE experte", "recurrence_rrule est facultative. Si remplie elle remplace frequency/interval/days. Ne pas préfixer par RRULE:."),
        ("Dates suivantes", "event_calendar.py reste l'agent dédié aux nouvelles dates, reports, annulations et clôtures."),
        ("Statuts", "pending = à traiter ; done = réussi ; error = à corriger/relancer ; skip = exemple ou ligne ignorée."),
    ]
    for r, (key, value) in enumerate(rows, 1):
        ws.cell(r, 1, key).font = Font(bold=True, color="FFFFFF")
        ws.cell(r, 1).fill = PatternFill("solid", fgColor="4472C4")
        ws.cell(r, 2, value).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 38
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Création et mise à jour des EventSeries, célébrations annuelles et événements récurrents.")
    parser.add_argument("--zone", required=True)
    parser.add_argument("--env", choices=["local", "staging", "prod"], default="local")
    parser.add_argument("--init", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--max", type=int, default=100)
    args = parser.parse_args()
    path = drive_workspace_root(args.env) / "izilife" / "events" / normalize_zone(args.zone) / FILE
    if args.init:
        example = {
            "object_type": "EVENT_SERIE", "string_id": "exemple-serie",
            "name": "Exemple — série d'événements", "description": "Description éditoriale complète.",
            "event_category_ref": "concert", "is_active": "non", "listable": "oui",
            "default_location_ref": "PLACE:exemple-lieu",
            "first_start_date": "2027-01-15", "first_start_hour": "19:00",
            "first_end_date": "2027-01-15", "first_end_hour": "23:00",
            "status": "skip", "notes": "Exemple non exécuté. Dupliquer la ligne et passer status à pending.",
        }
        validations = {
            "object_type": ["EVENT_SERIE", "ANNUAL_CELEBRATION", "RECURRENT_EVENT"],
            "is_active": ["oui", "non"], "listable": ["oui", "non"],
            "can_be_linked_by_others": ["oui", "non"], "is_internal": ["oui", "non"],
            "is_izilife_serie": ["oui", "non"], "is_exclusive_on_izilife": ["oui", "non"],
            "is_izilife_partnership_serie": ["oui", "non"], "is_izilife_organizer_serie": ["oui", "non"],
            "access_on_booking": ["oui", "non"],
            "recurrence_frequency": ["DAILY", "WEEKLY", "MONTHLY", "YEARLY"],
            "status": ["pending", "done", "error", "skip", "unmapped"],
        }
        create_workbook(path, SHEET, COLUMNS, example, validations)
        add_guide(path)
        return
    process(path, SHEET, args.env, "/scraper/agentUpsertEventStructure", payload, args.dry_run, args.max)


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from agent_excel_logger import run_logged
    run_logged("event_structures", "events", main)
