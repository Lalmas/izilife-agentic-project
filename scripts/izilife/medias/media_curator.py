#!/usr/bin/env python3
"""Curation et insertion directe de médias par scope iziLife."""
from __future__ import annotations

import argparse
from copy import copy
import hashlib
import json
import os
import re
import sys
import tempfile
from pathlib import Path

import openpyxl
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SCRIPT_ROOT = Path(__file__).resolve().parents[1]
if str(SCRIPT_ROOT) not in sys.path:
    sys.path.insert(0, str(SCRIPT_ROOT))

from objects.object_agent_common import (  # noqa: E402
    IZILIFE_ENVS, drive_workspace_root, load_env, normalize_env,
    normalize_zone, slugify,
)

MEDIA_EXT = {'.jpg', '.jpeg', '.png', '.webp', '.gif', '.mp4', '.mov', '.webm'}
SCOPE_TYPES = ['PLACE','SHOP','PAGE','EVENT','EVENT_SERIE','ANNUAL_CELEBRATION','EXPERIENCE','CITY','ADMINISTRATIVE_DIVISION','TERRACE']
GROUP_TYPES = ['SIMPLE','CAROUSEL','BEFORE','AFTER']
SOURCE_TYPES = ['INSTAGRAM','DIRECTORY']
YES = {'1','true','yes','oui','o','y'}
HEADERS = [
    'SCOPE_TYPE','SCOPE_REF','MEDIA_GROUP','SOURCE_TYPE','SOURCE_URLS',
    *[f'URL_{i}' for i in range(1, 11)],
    'DIRECTORY_NAME','TITLE','DESCRIPTION','PRINCIPAL','RIGHTS_CONFIRMED',
    'BELONGS_TO_SCOPE','STATUS','LAST_RESULT','OWNER_PAGE_STRING_ID','TAG_REFS',
]


def root_dir(zone: str, env: str) -> Path:
    return drive_workspace_root(env) / 'izilife' / 'medias' / normalize_zone(zone)


def default_file(zone: str, env: str) -> Path:
    return root_dir(zone, env) / 'media_curation.xlsx'


def truthy(value) -> bool:
    return str(value or '').strip().lower() in YES


def init_sheet(zone: str, env: str) -> None:
    root = root_dir(zone, env)
    (root / 'sources').mkdir(parents=True, exist_ok=True)
    (root / 'downloads').mkdir(parents=True, exist_ok=True)
    path = default_file(zone, env)
    if path.exists():
        wb = openpyxl.load_workbook(path)
        ws = wb['Medias'] if 'Medias' in wb.sheetnames else wb.active
        current = {str(cell.value or '').strip().upper() for cell in ws[1]}
        last_header_column = max((cell.column for cell in ws[1] if cell.value is not None), default=0)
        added = []
        for header in ('OWNER_PAGE_STRING_ID', 'TAG_REFS'):
            if header in current:
                continue
            column = last_header_column + 1
            source = ws.cell(1, max(1, column - 1))
            target = ws.cell(1, column, header)
            if source.has_style:
                target._style = copy(source._style)
            target.alignment = copy(source.alignment)
            ws.column_dimensions[target.column_letter].width = 32
            added.append(header)
            last_header_column = column
        guide_changed = False
        if 'Guide' in wb.sheetnames:
            guide = wb['Guide']
            existing_rules = {str(guide.cell(r, 1).value or '') for r in range(1, guide.max_row + 1)}
            for row_index in range(1, guide.max_row + 1):
                if str(guide.cell(row_index, 1).value or '').strip() == 'Droits':
                    new_rule = 'RIGHTS_CONFIRMED=OUI est obligatoire. BELONGS_TO_SCOPE=OUI rattache au scope ; NON exige OWNER_PAGE_STRING_ID.'
                    if guide.cell(row_index, 2).value != new_rule:
                        guide.cell(row_index, 2).value = new_rule
                        guide_changed = True
            for rule, value in [
                ('Propriétaire Page', 'Si BELONGS_TO_SCOPE=NON, OWNER_PAGE_STRING_ID est obligatoire (string_id de la Page).'),
                ('Tags', 'TAG_REFS accepte des références TYPE:string_id séparées par ;, par exemple CITY:lille;PLACE:soultraincafe.'),
            ]:
                if rule not in existing_rules:
                    guide.append([rule, value])
                    guide_changed = True
        if added or guide_changed:
            ws.auto_filter.ref = f'A1:{ws.cell(1, last_header_column).column_letter}500'
            wb.save(path)
            print(f'Classeur existant complété ({", ".join(added)}) : {path}')
        else:
            print(f'Classeur déjà à jour : {path}')
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Medias'
    ws.append(HEADERS)
    for index, header in enumerate(HEADERS, 1):
        cell = ws.cell(1, index)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 44 if header in {'SOURCE_URLS','DESCRIPTION','LAST_RESULT'} else 23
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{ws.cell(1, len(HEADERS)).column_letter}500'
    for column, values in [('A',SCOPE_TYPES),('C',GROUP_TYPES),('D',SOURCE_TYPES),('S',['OUI','NON']),('T',['OUI','NON']),('U',['OUI','NON'])]:
        dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"')
        ws.add_data_validation(dv)
        dv.add(f'{column}2:{column}500')
    guide = wb.create_sheet('Guide')
    guide.append(['Règle','Valeur'])
    guide.append(['Publication simple','MEDIA_GROUP=SIMPLE, une URL Instagram.'])
    guide.append(['Carrousel','MEDIA_GROUP=CAROUSEL. Toutes les images sont conservées dans un même sous-dossier ordonné.'])
    guide.append(['Plusieurs liens','SOURCE_URLS accepte un lien par ligne ou séparé par ;. URL_1 à URL_10 peuvent être utilisés en complément.'])
    guide.append(['Répertoire','SOURCE_TYPE=DIRECTORY et DIRECTORY_NAME relatif au dossier sources.'])
    guide.append(['Même cible','Répéter SCOPE_TYPE + SCOPE_REF sur autant de lignes que nécessaire.'])
    guide.append(['Droits','RIGHTS_CONFIRMED=OUI est obligatoire. BELONGS_TO_SCOPE=OUI rattache le média au scope ; NON exige OWNER_PAGE_STRING_ID.'])
    guide.append(['Propriétaire Page','Si BELONGS_TO_SCOPE=NON, OWNER_PAGE_STRING_ID est obligatoire (string_id de la Page).'])
    guide.append(['Tags','TAG_REFS accepte des références TYPE:string_id séparées par ;, par exemple CITY:lille;PLACE:soultraincafe.'])
    guide.append(['Principal','PRINCIPAL=OUI affecte seulement le premier fichier du lot.'])
    wb.save(path)
    print(f'Créé : {path}')
    print(f'Sources locales : {root / "sources"}')


def split_urls(value) -> list[str]:
    return [part.strip() for part in re.split(r'[;\n\r|]+', str(value or '')) if part.strip().startswith(('http://','https://'))]


def row_values(ws, row) -> dict:
    headers = [str(c.value or '').strip().upper() for c in ws[1]]
    return {headers[i]: row[i].value for i in range(min(len(headers), len(row))) if headers[i]}


def row_urls(values: dict) -> list[str]:
    urls = split_urls(values.get('SOURCE_URLS'))
    for index in range(1, 11):
        urls.extend(split_urls(values.get(f'URL_{index}')))
    return list(dict.fromkeys(urls))


def local_files(directory: Path) -> list[Path]:
    if not directory.is_dir():
        raise RuntimeError(f'Répertoire introuvable : {directory}')
    files = sorted(p for p in directory.rglob('*') if p.is_file() and p.suffix.lower() in MEDIA_EXT)
    if not files:
        raise RuntimeError(f'Aucun média supporté dans : {directory}')
    return files


def dedupe_files(files: list[Path]) -> list[Path]:
    seen, result = set(), []
    for path in files:
        digest = hashlib.sha1(path.read_bytes()).hexdigest()
        if digest not in seen:
            seen.add(digest)
            result.append(path)
    return result


def prepare_upload(path: Path, temporary_directory: Path) -> Path:
    """Resize oversized still images for the API without altering source files."""
    if path.suffix.lower() not in {'.jpg', '.jpeg', '.png', '.webp'}:
        return path
    try:
        from PIL import Image, ImageOps
    except ImportError as exc:
        raise RuntimeError('Pillow est requis pour adapter les dimensions avant upload.') from exc

    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source)
        if image.width <= 2000 and image.height <= 2000 and path.stat().st_size <= 5 * 1024 * 1024:
            return path
        image.thumbnail((2000, 2000), Image.Resampling.LANCZOS)
        has_alpha = image.mode in {'RGBA', 'LA'} or 'transparency' in image.info
        destination = temporary_directory / (path.stem + ('_upload.png' if has_alpha else '_upload.jpg'))
        if has_alpha:
            image.convert('RGBA').save(destination, format='PNG', optimize=True)
        else:
            image.convert('RGB').save(destination, format='JPEG', quality=88, optimize=True)
    if destination.stat().st_size > 5 * 1024 * 1024:
        raise RuntimeError(f'Image trop lourde après redimensionnement : {path.name}')
    return destination


def upload(path: Path, values: dict, principal: bool, env: str) -> dict:
    cfg = IZILIFE_ENVS[env]
    token = os.getenv('IZILIFE_AGENT_TOKEN','').strip()
    if not token:
        raise RuntimeError('IZILIFE_AGENT_TOKEN manquant')
    data = {
        'scope_type': str(values.get('SCOPE_TYPE') or '').strip().upper(),
        'scope_ref': str(values.get('SCOPE_REF') or '').strip(),
        'title': str(values.get('TITLE') or path.stem).strip(),
        'description': str(values.get('DESCRIPTION') or '').strip(),
        'is_principal': 'oui' if principal else 'non',
        'rights_confirmed': 'oui' if truthy(values.get('RIGHTS_CONFIRMED')) else 'non',
        'belongs_to_scope': 'oui' if truthy(values.get('BELONGS_TO_SCOPE')) else 'non',
        'owner_page_string_id': str(values.get('OWNER_PAGE_STRING_ID') or '').strip(),
        'tag_refs': json.dumps([item.strip() for item in re.split(r'[;\n\r|]+', str(values.get('TAG_REFS') or '')) if item.strip()], ensure_ascii=False),
    }
    with path.open('rb') as handle:
        response = requests.post(
            cfg['base_url'].rstrip('/') + '/scraper/agentAddScopedMedia',
            data=data, files={'media': (path.name, handle)},
            headers={'X-Agent-Token': token}, verify=cfg.get('verify_ssl', True), timeout=240,
        )
    try:
        result = response.json()
    except Exception:
        result = {'success':False,'error':f'HTTP {response.status_code}: {response.text[:500]}'}
    return result


def set_result(ws, row_number: int, status: str, message: str) -> None:
    headers = [str(c.value or '').strip().upper() for c in ws[1]]
    ws.cell(row_number, headers.index('STATUS') + 1, status)
    ws.cell(row_number, headers.index('LAST_RESULT') + 1, message[:32000])


def run(zone: str, env: str, file: str | None, dry_run: bool, maximum: int) -> None:
    load_env()
    env = normalize_env(env)
    path = Path(file) if file else default_file(zone, env)
    if not path.exists():
        raise SystemExit(f'Fichier introuvable : {path}. Lance --init.')
    wb = openpyxl.load_workbook(path)
    ws = wb['Medias'] if 'Medias' in wb.sheetnames else wb.active
    pending = []
    for row in ws.iter_rows(min_row=2):
        values = row_values(ws, row)
        status = str(values.get('STATUS') or '').strip().lower()
        if status in {'done','skip','skipped'}:
            continue
        if not (str(values.get('SCOPE_REF') or '').strip() or str(values.get('OWNER_PAGE_STRING_ID') or '').strip()):
            continue
        pending.append((row[0].row, values))
        if maximum > 0 and len(pending) >= maximum:
            break
    if not pending:
        print('Aucune ligne à traiter.')
        return
    for _, values in pending:
        if not truthy(values.get('RIGHTS_CONFIRMED')):
            raise SystemExit('Chaque ligne doit confirmer RIGHTS_CONFIRMED=OUI avant lancement.')
        if str(values.get('BELONGS_TO_SCOPE') or '').strip().lower() not in {'oui','non','1','0','true','false','yes','no','o','n','y'}:
            raise SystemExit('BELONGS_TO_SCOPE doit être renseigné explicitement à OUI ou NON.')
        if truthy(values.get('BELONGS_TO_SCOPE')) and (not str(values.get('SCOPE_TYPE') or '').strip() or not str(values.get('SCOPE_REF') or '').strip()):
            raise SystemExit('BELONGS_TO_SCOPE=OUI exige SCOPE_TYPE et SCOPE_REF.')
        if not truthy(values.get('BELONGS_TO_SCOPE')) and not str(values.get('OWNER_PAGE_STRING_ID') or '').strip():
            raise SystemExit('BELONGS_TO_SCOPE=NON exige OWNER_PAGE_STRING_ID.')
    if dry_run:
        for row_number, values in pending:
            print('[DRY RUN]', row_number, values.get('SCOPE_TYPE'), values.get('SCOPE_REF'), 'owner_page=', values.get('OWNER_PAGE_STRING_ID'), 'tags=', values.get('TAG_REFS'), values.get('MEDIA_GROUP'), values.get('SOURCE_TYPE'), row_urls(values) or values.get('DIRECTORY_NAME'))
        return

    instagram_rows = [(n,v) for n,v in pending if str(v.get('SOURCE_TYPE') or '').strip().upper() == 'INSTAGRAM']
    browser = context = page_browser = playwright = None
    if instagram_rows:
        from playwright.sync_api import sync_playwright
        from media_downloader import apply_stealth, open_igram_session
        playwright = sync_playwright().start()
        try:
            browser = playwright.chromium.launch(channel='chrome', headless=False, args=['--disable-blink-features=AutomationControlled'])
        except Exception:
            browser = playwright.chromium.launch(headless=False, args=['--disable-blink-features=AutomationControlled'])
        context = browser.new_context(locale='fr-FR')
        page_browser = context.new_page()
        apply_stealth(page_browser)
        open_igram_session(page_browser)

    stats = {'done':0,'created':0,'skipped':0,'errors':0}
    try:
        for row_number, values in pending:
            try:
                scope_type = str(values.get('SCOPE_TYPE') or '').strip().upper()
                scope_ref = str(values.get('SCOPE_REF') or '').strip()
                group = str(values.get('MEDIA_GROUP') or 'SIMPLE').strip().upper()
                source_type = str(values.get('SOURCE_TYPE') or '').strip().upper()
                if (scope_ref and scope_type not in SCOPE_TYPES) or group not in GROUP_TYPES or source_type not in SOURCE_TYPES:
                    raise RuntimeError('SCOPE_TYPE, MEDIA_GROUP ou SOURCE_TYPE invalide.')
                if truthy(values.get('BELONGS_TO_SCOPE')) and (scope_type not in SCOPE_TYPES or not scope_ref):
                    raise RuntimeError('Un scope valide est obligatoire quand BELONGS_TO_SCOPE=OUI.')
                if source_type == 'DIRECTORY':
                    directory_name = str(values.get('DIRECTORY_NAME') or '').strip()
                    directory = (root_dir(zone, env) / 'sources' / directory_name).resolve()
                    source_root = (root_dir(zone, env) / 'sources').resolve()
                    if source_root not in directory.parents and directory != source_root:
                        raise RuntimeError('DIRECTORY_NAME doit rester dans le dossier sources.')
                    files = local_files(directory)
                else:
                    from media_downloader import download_one_url
                    urls = row_urls(values)
                    if not urls:
                        raise RuntimeError('Aucune URL Instagram.')
                    download_type = scope_type or 'page'
                    download_ref = scope_ref or str(values.get('OWNER_PAGE_STRING_ID') or '')
                    lot = root_dir(zone, env) / 'downloads' / download_type.lower() / slugify(download_ref) / f'{row_number:04d}_{group.lower()}'
                    lot.mkdir(parents=True, exist_ok=True)
                    files = []
                    media_type = 'carousel' if group == 'CAROUSEL' else 'photo'
                    for url in urls:
                        files.extend(download_one_url(page_browser, url, media_type, lot, lot))
                files = dedupe_files(files)
                created = skipped = 0
                errors = []
                with tempfile.TemporaryDirectory(prefix='izilife_media_upload_') as temporary_directory:
                    temporary_path = Path(temporary_directory)
                    for index, media in enumerate(files):
                        upload_path = prepare_upload(media, temporary_path)
                        result = upload(upload_path, values, truthy(values.get('PRINCIPAL')) and index == 0, env)
                        if result.get('success'):
                            if result.get('created'): created += 1
                            else: skipped += 1
                        else:
                            errors.append(f'{media.name}: {result.get("error", result)}')
                if errors:
                    raise RuntimeError(' | '.join(errors))
                message = json.dumps({'files':len(files),'created':created,'skipped':skipped}, ensure_ascii=False)
                set_result(ws, row_number, 'done', message)
                wb.save(path)
                stats['done'] += 1; stats['created'] += created; stats['skipped'] += skipped
                print(f'OK ligne {row_number}: {message}')
            except Exception as exc:
                stats['errors'] += 1
                set_result(ws, row_number, 'error', str(exc))
                wb.save(path)
                print(f'ERREUR ligne {row_number}: {exc}')
    finally:
        if context: context.close()
        if browser: browser.close()
        if playwright: playwright.stop()
    print('Résultat:', json.dumps(stats, ensure_ascii=False))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--zone', required=True)
    parser.add_argument('--env', default='local', choices=['local','staging','prod'])
    parser.add_argument('--file')
    parser.add_argument('--init', action='store_true')
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--max', type=int, default=500)
    args = parser.parse_args()
    init_sheet(args.zone, normalize_env(args.env)) if args.init else run(args.zone, args.env, args.file, args.dry_run, args.max)


if __name__ == '__main__':
    from agent_excel_logger import run_logged
    run_logged('medias', 'media-curation', main)
