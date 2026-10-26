#!/usr/bin/env python3
"""Enrichisseur Product/Service/EAP depuis une liste explicite de liens."""
from __future__ import annotations

import argparse, json, re, sys, time
from pathlib import Path
from urllib.parse import quote, urlparse
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from playwright.sync_api import sync_playwright

SCRIPT_ROOT = Path(__file__).resolve().parents[1]
if str(SCRIPT_ROOT) not in sys.path: sys.path.insert(0, str(SCRIPT_ROOT))
from objects.object_agent_common import drive_workspace_root, izilife_post, load_env, normalize_env, normalize_zone

HEADERS = ['SCOPE_TYPE','SCOPE_REF','IMAGE_DIRECTORY',*[f'URL_{i}' for i in range(1,11)],'NOTES','LLM_INSTRUCTIONS','STATUS','LAST_RESULT']
PROCESSABLE_STATUSES = {'','pending','retry'}
STATUS_VALUES = ['pending','done','review','retry','error','skip']

def root_dir(zone, env):
    return drive_workspace_root(env)/'izilife'/'places'/normalize_zone(zone)

def workbook_file(zone, env): return root_dir(zone,env)/'enrichisseur_contenus.xlsx'

def configure_sources_sheet(ws):
    """Aligne un nouveau classeur ou un classeur historique sur le contrat courant."""
    changed=False
    headers=[str(c.value or '').strip().upper() for c in ws[1]]
    for required in HEADERS:
        if required in headers: continue
        if required == 'LLM_INSTRUCTIONS' and 'STATUS' in headers:
            column=headers.index('STATUS')+1
            ws.insert_cols(column)
            ws.cell(1,column,required)
        else:
            ws.cell(1,ws.max_column+1,required)
        headers=[str(c.value or '').strip().upper() for c in ws[1]]
        changed=True

    for i,h in enumerate(headers,1):
        c=ws.cell(1,i); c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78'); c.alignment=Alignment(horizontal='center')
        ws.column_dimensions[c.column_letter].width=65 if h in {'LLM_INSTRUCTIONS','LAST_RESULT'} else (48 if h.startswith('URL_') else 28)
    ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:{ws.cell(1,len(headers)).column_letter}500'

    # Supprime seulement les anciennes listes gérées par ce script avant de les recréer.
    kept=[]
    for validation in ws.data_validations.dataValidation:
        formula=str(validation.formula1 or '').lower()
        if ('place,shop' in formula) or ('pending' in formula and 'done' in formula and 'skip' in formula):
            changed=True
            continue
        kept.append(validation)
    ws.data_validations.dataValidation=kept
    scope_col=headers.index('SCOPE_TYPE')+1; status_col=headers.index('STATUS')+1
    dv=DataValidation(type='list',formula1='"PLACE,SHOP"'); ws.add_data_validation(dv); dv.add(f'{ws.cell(1,scope_col).column_letter}2:{ws.cell(1,scope_col).column_letter}500')
    dv2=DataValidation(type='list',formula1='"'+','.join(STATUS_VALUES)+'"'); ws.add_data_validation(dv2); dv2.add(f'{ws.cell(1,status_col).column_letter}2:{ws.cell(1,status_col).column_letter}500')
    return changed

def configure_guide_sheet(wb):
    guide=wb['Guide'] if 'Guide' in wb.sheetnames else wb.create_sheet('Guide')
    rows=[
        ['Règle','Valeur'],
        ['Une ligne','Une cible scope_type + scope_ref.'],
        ['Photos menu / flyer','Créer un sous-répertoire à côté de ce classeur et écrire uniquement son nom dans IMAGE_DIRECTORY. Toutes ses images seront analysées.'],
        ['Liens','Une seule URL par colonne, de URL_1 à URL_10. Ne jamais regrouper plusieurs URL dans une cellule.'],
        ['Navigation','Seuls les liens renseignés sont ouverts. Aucun lien interne n’est suivi.'],
        ['Instructions LLM','LLM_INSTRUCTIONS précise une correction métier à appliquer lors de la prochaine analyse.'],
        ['Statuts','pending = à traiter; done = terminé; review = décision métier attendue; retry = relancer avec les instructions; error = échec technique; skip = ignorer.'],
        ['Relance','Seuls pending, retry et une cellule STATUS vide sont traités automatiquement.'],
        ['Résultat','Chaque lien ou image est analysé séparément puis les objets certains sont insérés par upsert. Les avertissements sont conservés dans LAST_RESULT.'],
    ]
    guide.delete_rows(1,guide.max_row)
    for row in rows: guide.append(row)
    guide.column_dimensions['A'].width=24; guide.column_dimensions['B'].width=115
    for cell in guide[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78')

def init_sheet(zone, env):
    path=workbook_file(zone,env); path.parent.mkdir(parents=True,exist_ok=True)
    if path.exists():
        wb=openpyxl.load_workbook(path); ws=wb['Sources'] if 'Sources' in wb.sheetnames else wb.active
        configure_sources_sheet(ws); configure_guide_sheet(wb); wb.save(path); print(f'Mis à jour : {path}'); return
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='Sources'; ws.append(HEADERS)
    configure_sources_sheet(ws); configure_guide_sheet(wb)
    wb.save(path); print(f'Créé : {path}')

def split_urls(value): return [v.strip() for v in re.split(r'[;\n\r|]+',str(value or '')) if v.strip().startswith(('http://','https://'))]

def extract_urls(values):
    urls=[]
    for i in range(1,11):
        value=str(values.get(f'URL_{i}') or '').strip()
        if not value: continue
        parsed=split_urls(value)
        if len(parsed) != 1 or parsed[0] != value:
            raise RuntimeError(f'URL_{i} doit contenir une seule URL complète.')
        urls.append(value)
    return list(dict.fromkeys(urls))

IMAGE_EXTENSIONS={'.jpg','.jpeg','.png','.webp','.gif','.bmp'}

def extract_images(values, workbook_path):
    directory_text=str(values.get('IMAGE_DIRECTORY') or '').strip()
    if not directory_text: return []
    directory=Path(directory_text).expanduser()
    if not directory.is_absolute(): directory=Path(workbook_path).parent/directory
    candidates=list(directory.iterdir()) if directory.is_dir() else []
    images=[]
    for candidate in candidates:
        path=candidate.resolve()
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS and path not in images: images.append(path)
    if not images:
        raise RuntimeError(f'Aucune image trouvée dans le répertoire {directory}.')
    return images

def compact_review_items(items):
    compact=[]
    for entry in items:
        if not isinstance(entry,dict): continue
        item=entry.get('item') if isinstance(entry.get('item'),dict) else {}
        compact.append({
            'kind':entry.get('kind'),
            'name':item.get('name'),
            'category_name':item.get('category_name'),
            'element_menu':item.get('element_menu'),
        })
    return compact

def dedupe_extracted_items(items):
    """Supprime les répétitions créées par le chevauchement des blocs source.

    La clé suit aussi l'identité réellement utilisée par les upserts. Un même
    Service/Equipment/Animation extrait dans plusieurs blocs ne doit jamais être
    recréé parce que le LLM a varié la casse, la catégorie ou un champ tarifaire.
    Les Product restent distincts par carte et les EAP par variante tarifaire.
    En cas de doublon, la représentation la plus renseignée est conservée.
    """
    unique={}
    for entry in items:
        if not isinstance(entry,dict) or not isinstance(entry.get('item'),dict): continue
        kind=str(entry.get('kind') or '').strip().casefold(); item=entry['item']
        menu=item.get('element_menu') if isinstance(item.get('element_menu'),dict) else {}
        category_path=item.get('category_path') if isinstance(item.get('category_path'),list) else []
        def norm(value): return re.sub(r'\s+',' ',str(value or '')).strip().casefold()
        name=norm(item.get('name')); menu_ref=norm(menu.get('ref') or menu.get('name'))
        if kind=='product': key=(kind,name,menu_ref)
        elif kind=='eap': key=(kind,name,norm(item.get('price')),norm(item.get('duration_minutes')),norm(item.get('validity_period')))
        else: key=(kind,name)
        if not key[0] or not key[1]: continue
        previous=unique.get(key)
        if previous is None or len(json.dumps(entry,ensure_ascii=False,sort_keys=True))>len(json.dumps(previous,ensure_ascii=False,sort_keys=True)):
            unique[key]=entry
    return list(unique.values())

def consolidate_element_menus(items):
    """Réconcilie les noms de cartes inventés différemment par des blocs voisins.

    Un libellé explicitement incertain n'est jamais autorisé à créer une nouvelle
    carte. S'il contient le nom stable d'une carte observée ailleurs, il est
    rattaché à celle-ci; sinon le produit reste au niveau du scope et un warning
    est remonté. Cette règle ne dépend d'aucun établissement particulier.
    """
    def norm(value):
        value=re.sub(r'[^\w\s-]+',' ',str(value or '').casefold(),flags=re.UNICODE)
        return re.sub(r'\s+',' ',value).strip()
    uncertain=re.compile(r'(\?|\b(?:probable|suppos[ée]e?|incertain|[àa] confirmer)\b)',re.I)
    stable={}
    for entry in items:
        item=entry.get('item') if isinstance(entry,dict) else None
        menu=item.get('element_menu') if isinstance(item,dict) and isinstance(item.get('element_menu'),dict) else None
        name=str((menu or {}).get('name') or (menu or {}).get('public_label') or '').strip()
        if name and not uncertain.search(name): stable.setdefault(norm(name),menu)
    warnings=[]
    for entry in items:
        item=entry.get('item') if isinstance(entry,dict) else None
        menu=item.get('element_menu') if isinstance(item,dict) and isinstance(item.get('element_menu'),dict) else None
        name=str((menu or {}).get('name') or (menu or {}).get('public_label') or '').strip()
        if not name or not uncertain.search(name): continue
        candidate=norm(name)
        matches=[(key,value) for key,value in stable.items() if key and (key in candidate or candidate in key)]
        if matches:
            _,canonical=max(matches,key=lambda pair:len(pair[0]))
            item['element_menu']=dict(canonical)
            warnings.append(f'ElementMenu incertain « {name} » rattaché à « {canonical.get("name") or canonical.get("public_label")} ».')
        else:
            item['element_menu']=None
            warnings.append(f'ElementMenu incertain « {name} » ignoré; produit conservé sur le scope.')
    # Lorsqu'une page expose plusieurs cartes et qu'une seule est explicitement
    # la carte générale, les produits sans contexte répété par les blocs suivants
    # héritent de cette carte. Les cartes spécialisées (brunch, midi, soir, etc.)
    # gardent uniquement les produits que la source leur attribue explicitement.
    defaults=[]
    for key,menu in stable.items():
        presentation=str(menu.get('presentation_type') or '').upper()
        if key in {'carte','menu','catalogue','card'} or (presentation in {'CARD','CATALOG'} and not menu.get('is_for_brunch')):
            defaults.append(menu)
    if len(stable)>=2 and len(defaults)==1:
        default_menu=defaults[0]
        inherited=0
        for entry in items:
            item=entry.get('item') if isinstance(entry,dict) else None
            if str(entry.get('kind') or '').casefold()!='product' or not isinstance(item,dict): continue
            if not isinstance(item.get('element_menu'),dict):
                item['element_menu']=dict(default_menu); inherited+=1
        if inherited:
            warnings.append(f'{inherited} produit(s) sans contexte de bloc rattaché(s) à la carte générale « {default_menu.get("name") or default_menu.get("public_label")} ».')
    return items,warnings

def ingest_menu_image(image_path,scope,ref,env):
    mime={'.png':'image/png','.webp':'image/webp','.gif':'image/gif','.bmp':'image/bmp'}.get(image_path.suffix.lower(),'image/jpeg')
    with image_path.open('rb') as stream:
        return izilife_post('/scraper/agentIngestProductImage',{
            'scope_type':scope,'scope_ref':ref,'source_name':image_path.name,
        },env,timeout=300,files={'image':(image_path.name,stream,mime)})

def dismiss_cookie_consent(page):
    """Ferme les CMP courantes, en privilégiant le refus des cookies facultatifs."""
    selectors = [
        '#onetrust-reject-all-handler', '#CybotCookiebotDialogBodyButtonDecline',
        '[data-testid="uc-deny-all-button"]', '[data-testid="cookie-policy-manage-dialog-decline"]',
        '#didomi-notice-disagree-button', '.didomi-continue-without-agreeing',
        '[data-qa="oil-NoButton"]', '#axeptio_btn_dismiss', '.qc-cmp2-summary-buttons button:has-text("Refuser")',
        '.truste-consent-button', '#tarteaucitronAllDenied2', '.iubenda-cs-reject-btn',
        'button:has-text("Tout refuser")', 'button:has-text("Refuser tout")',
        'button:has-text("Continuer sans accepter")', 'button:has-text("Reject all")',
        'button:has-text("Refuse all")', '[aria-label*="Tout refuser" i]',
        '[aria-label*="Reject all" i]',
        '#onetrust-accept-btn-handler', '#CybotCookiebotDialogBodyLevelButtonLevelOptinAllowAll',
        '#didomi-notice-agree-button', '#axeptio_btn_acceptAll', '#tarteaucitronAllAllowed', '.iubenda-cs-accept-btn',
        '[data-testid="uc-accept-all-button"]', 'button:has-text("Tout accepter")',
        'button:has-text("Accepter tout")', 'button:has-text("Accept all")',
        '[aria-label*="Tout accepter" i]', '[aria-label*="Accept all" i]',
    ]
    for _ in range(4):
        clicked = False
        for target in [page, *page.frames]:
            for selector in selectors:
                try:
                    button = target.locator(selector).first
                    if button.is_visible(timeout=250):
                        button.click(force=True, timeout=2500)
                        clicked = True
                        break
                except Exception:
                    pass
            if clicked: break
        if not clicked: break
        page.wait_for_timeout(1200)
    try: page.keyboard.press('Escape')
    except Exception: pass
    # Dernier recours : fermer uniquement un dialogue/overlay réellement
    # bloquant et identifiable comme consentement, jamais un élément métier.
    cookie_words=re.compile(r'cookie|consentement|confidentialit|privacy|traceur',re.I)
    for target in [page,*page.frames]:
        for modal_selector in ['[role="dialog"]','.modal.show','.popup.is-visible','.overlay.active','[aria-modal="true"]']:
            try:
                for index in range(min(target.locator(modal_selector).count(),8)):
                    modal=target.locator(modal_selector).nth(index)
                    if not modal.is_visible(timeout=150): continue
                    text=modal.inner_text(timeout=500)
                    if not cookie_words.search(text or ''): continue
                    close=modal.locator('button[aria-label*="fermer" i],button[aria-label*="close" i],[title*="fermer" i],[title*="close" i],button.close,.modal-close,.close').first
                    if close.is_visible(timeout=200): close.click(force=True,timeout=1500)
            except Exception: pass
    try:
        page.evaluate("""() => { document.documentElement.style.overflow='auto'; document.body.style.overflow='auto'; }""")
    except Exception: pass

def scrape_exact_url(page,url):
    try:
        page.goto(url,wait_until='domcontentloaded',timeout=45000)
    except Exception as first_error:
        # Certains moteurs de réservation conservent indéfiniment une requête
        # ouverte : le DOM utile est pourtant déjà disponible. Si rien n'est
        # exploitable, retenter en n'attendant que la réponse initiale.
        try:
            visible=page.locator('body').inner_text(timeout=1500).strip()
        except Exception:
            visible=''
        if len(visible)<40:
            try:
                page.goto(url,wait_until='commit',timeout=60000)
                page.wait_for_timeout(5000)
            except Exception as retry_error:
                raise RuntimeError(f'Navigation impossible après nouvelle tentative: {retry_error}') from first_error
    try: page.wait_for_load_state('networkidle',timeout=8000)
    except Exception: pass
    dismiss_cookie_consent(page)
    # Capturer chaque état pendant le scroll : les listes virtualisées peuvent
    # retirer du DOM les produits déjà dépassés quand on atteint le bas.
    snapshots=[]; image_contexts=[]; interactive_contexts=[]; previous_signature=None; stagnant=0
    host=(urlparse(url).hostname or '').lower()
    max_scrolls=15 if 'halluneed.com' in host else 30
    frames_to_read=[page.main_frame] if 'halluneed.com' in host else page.frames[:8]
    for _ in range(max_scrolls):
        texts=[]
        # Les pages marketing peuvent charger des dizaines d'iframes publicitaires.
        # Leur lecture séquentielle rendait le test quasi bloquant : le document
        # principal et les premiers iframes métier suffisent ici.
        for frame in frames_to_read:
            try:
                value=frame.locator('body').inner_text(timeout=500).strip()
                if value and value not in texts: texts.append(value)
            except Exception: pass
        text='\n\n'.join(texts)
        if len(text)>=40 and text not in snapshots: snapshots.append(text)
        try:
            images=page.locator('img').evaluate_all(r"""imgs => imgs.map(img => {
                const owner=img.closest('article,li,[role="listitem"],[data-testid*="item"],.card') || img.parentElement;
                const context=((owner?.innerText||'').trim() || img.alt || '').slice(0,800);
                return {url:img.currentSrc||img.src||'',w:img.naturalWidth||0,h:img.naturalHeight||0,text:context};
            }).filter(x => /^https?:\/\//i.test(x.url) && x.w>=80 && x.h>=80 && x.text.length>1)""")
            for image in images:
                line=f"IMAGE_URL: {image['url']}\nIMAGE_CONTEXT: {image['text']}"
                if line not in image_contexts: image_contexts.append(line)
        except Exception: pass
        try:
            signature=page.evaluate("""() => {
                const all=[document.scrollingElement,...document.querySelectorAll('*')].filter(Boolean);
                const scrollables=all.filter(el => el.scrollHeight > el.clientHeight + 80);
                scrollables.forEach(el => { el.scrollTop=el.scrollHeight; });
                window.scrollTo(0,document.documentElement.scrollHeight);
                return `${document.body.innerText.length}|${document.images.length}|${scrollables.map(x=>x.scrollHeight).join(',')}`;
            }""")
        except Exception: signature=str(len(text))
        stagnant=stagnant+1 if signature==previous_signature else 0; previous_signature=signature
        if stagnant>=4: break
        page.wait_for_timeout(1200)
        dismiss_cookie_consent(page)
    # Capturer aussi les états d'onglets/carrousels éditoriaux. Sur plusieurs
    # sites, les capacités et quantités ne sont présentes dans le DOM qu'après
    # activation de chaque thème (karaoké, bowling, missions, etc.).
    explored=set()
    interactive_seen={re.sub(r'\s+',' ',line).strip().casefold() for snapshot in snapshots for line in snapshot.splitlines() if line.strip()}
    for selector in [
        '[role="tab"]','[data-bs-toggle="tab"]','[data-toggle="tab"]','.swiper-menu-item',
        '[class*="menu-picker"] button','button[class*="menu-picker"]',
        '[class*="menu-switch"] button','[class*="menu-tabs"] button','[class*="menu-selector"] button',
    ]:
        try: candidates=page.locator(selector)
        except Exception: continue
        for index in range(min(candidates.count(),24)):
            candidate=candidates.nth(index)
            try:
                label=re.sub(r'\s+',' ',candidate.inner_text(timeout=300)).strip()
                key=(selector,label.casefold())
                if not label or key in explored or not candidate.is_visible(timeout=200): continue
                explored.add(key); candidate.click(force=True,timeout=2000); page.wait_for_timeout(500)
                state=page.locator('body').inner_text(timeout=1200).strip()
                if len(state)>=40 and state not in snapshots: snapshots.append(state)
                state_lines=[re.sub(r'\s+',' ',line).strip() for line in state.splitlines() if line.strip()]
                delta_lines=[]
                for line in state_lines:
                    normalized=line.casefold()
                    if normalized not in interactive_seen: delta_lines.append(line)
                    interactive_seen.add(normalized)
                if delta_lines:
                    interactive_contexts.append('INTERACTIVE_STATE: '+label+'\n'+'\n'.join(delta_lines))
                focused=[]
                for line_index,line in enumerate(state_lines):
                    if re.match(r'^\d+\s+(?:salles?|pistes?|box|machines?|bornes?|tables?)$',line,re.I):
                        focused.extend(state_lines[max(0,line_index-2):min(len(state_lines),line_index+5)])
                if focused:
                    interactive_contexts.append('INTERACTIVE_STATE: '+label+'\n'+'\n'.join(focused))
            except Exception: pass
    # Conserver l'ordre et le niveau des titres avant d'aplatir le texte. Cette
    # vue structurelle sert à reconstruire catégories et sous-catégories sur les
    # pages longues, indépendamment de leur découpage en blocs.
    structure_lines=[]
    try:
        structure=page.locator('h1,h2,h3,h4,[role="tab"],nav a').evaluate_all("""nodes => nodes.map((node,index) => ({
            index,
            tag:(node.getAttribute('role')==='tab'?'TAB':node.tagName),
            text:(node.innerText||node.textContent||'').replace(/\\s+/g,' ').trim()
        })).filter(x => x.text && x.text.length <= 160)""")
        seen_structure=set()
        for node in structure:
            line=f"{node['tag']}: {node['text']}"
            key=line.casefold()
            if key not in seen_structure: seen_structure.add(key); structure_lines.append(line)
    except Exception: pass
    # Les snapshots successifs répètent presque toute la page. Dédupliquer les
    # lignes évite d'envoyer 5 à 10 fois le même tarif au LLM et accélère
    # fortement les sites vitrines riches comme Hall U Need.
    unique_lines=[]; seen_lines=set()
    for snapshot in snapshots:
        for raw_line in snapshot.splitlines():
            line=re.sub(r'\s+',' ',raw_line).strip()
            key=line.casefold()
            if line and key not in seen_lines: seen_lines.add(key); unique_lines.append(line)
    content=('=== STRUCTURE DOM ORDONNÉE ===\n'+'\n'.join(structure_lines)+'\n\n=== TEXTE VISIBLE ===\n' if structure_lines else '')+'\n'.join(unique_lines)
    if interactive_contexts: content+='\n\n=== ÉTATS INTERACTIFS DÉTECTÉS ===\n'+'\n\n'.join(interactive_contexts)
    if image_contexts: content+='\n\n=== IMAGES PRODUITS DÉTECTÉES ===\n'+'\n\n'.join(image_contexts[:40])
    if len(content)<40: raise RuntimeError('Aucun texte visible après gestion des cookies, lecture des iframes et scroll exhaustif.')
    return content[:250000]

def content_chunks(content, size=2500, overlap=500):
    if len(content)<=size: return [content]
    chunks=[]; start=0
    while start<len(content):
        end=min(len(content),start+size); chunks.append(content[start:end])
        if end>=len(content): break
        start=max(start+1,end-overlap)
    return chunks

def parse_price(value):
    """Convertit les montants Dokaa ('12,90 €') sans inventer de prix."""
    match=re.search(r'-?\d+(?:[.,]\d+)?',str(value or '').replace('\u202f',' ').replace('\xa0',' '))
    return float(match.group(0).replace(',','.')) if match else None

def normalize_quantity_label(label):
    """Uniformise seulement les quantités certaines, sans réécrire le métier."""
    text=str(label or '').strip()
    match=re.match(r'^(\d+(?:[.,]\d+)?)\s*(ml|cl|l)$',text,re.IGNORECASE)
    if match: return f'{match.group(1)} {match.group(2).lower()}'
    return text

def infer_product_family(name, category):
    """Propose une famille canonique ; le BO crée un candidat PENDING si inconnue."""
    haystack=f'{category} {name}'.casefold()
    families=(
        ('mocktail','Mocktail'),('cocktail','Cocktail'),('bière','Bière'),('biere','Bière'),
        ('vin ','Vin'),('champagne','Champagne et vin effervescent'),('espresso','Espresso'),
        ('cappuccino','Cappuccino'),('latte','Latte'),('café','Café'),('cafe','Café'),
        ('thé','Thé'),('the ','Thé'),('smoothie','Smoothie'),('limonade','Limonade'),
        ('cola','Cola'),('soda','Soda'),('jus','Jus de fruit'),('eau','Eau'),
        ('pizza','Pizza'),('burger','Burger'),('sandwich','Sandwich'),('tacos','Tacos'),
        ('kebab','Kebab'),('wrap','Wrap'),('poké','Poké'),('poke','Poké'),('sushi','Sushi'),
        ('ramen','Ramen'),('pâtes','Pâtes'),('pates','Pâtes'),('salade','Salade composée'),
        ('soupe','Soupe'),('tapas','Tapas'),('planche','Planche à partager'),
        ('pain','Pain'),('croissant','Viennoiserie'),('gâteau','Gâteau'),('gateau','Gâteau'),
        ('glace','Glace et sorbet'),('dessert','Gâteau'),('fromage','Fromage'),
    )
    for token,family in families:
        if token in haystack: return family
    return category.strip() or name.strip()

def scrape_dokaa_products(page,url):
    """Lit le JSON natif Dokaa : aucune interprétation LLM des catégories/variantes."""
    parsed=urlparse(url); slug=parsed.path.rstrip('/').split('/')[-1]
    if not slug: raise RuntimeError('Identifiant de menu Dokaa introuvable dans l’URL.')
    api_path=f'/api/qr/uri/{quote(slug)}'
    with page.expect_response(lambda response: api_path in response.url,timeout=45000) as response_info:
        page.goto(url,wait_until='domcontentloaded',timeout=45000)
    response=response_info.value
    if not response.ok: raise RuntimeError(f'API Dokaa HTTP {response.status}: {response.url}')
    data=response.json(); languages=(data.get('data') or {}).get('languages') or {}
    language=languages.get('fr') or next(iter(languages.values()),{})
    products=[]
    for section in language.get('sections') or []:
        category=str(section.get('name') or '').strip()
        if not category: continue
        for item in section.get('items') or []:
            if item.get('disabled') is True: continue
            name=str(item.get('name') or '').strip()
            # « NEW - » est un marqueur éditorial Dokaa, pas une partie du
            # nom métier. La mise en avant sera gérée par les highlights iziLife.
            name=re.sub(r'^\s*new\s*!?\s*[-–—:]\s*','',name,flags=re.IGNORECASE).strip()
            if not name: continue
            item_id=str(item.get('id') or '').strip()
            descriptions=[]
            for field in ('translatedName','description'):
                value=str(item.get(field) or '').strip()
                if value and value not in descriptions: descriptions.append(value)
            descriptions.append('Source: dokaa')
            product={
                'name':name,
                'description':' — '.join(descriptions),
                'category_name':category,
                'category_path':[category],
                'product_family_name':infer_product_family(name,category),
                'currency_code':'EUR',
                'raw_source_url':f'{url}#dokaa:{item_id}' if item_id else url,
            }
            image=str(item.get('image') or '').strip()
            if image:
                # Toujours importer l'original Dokaa. La variante /medium est une
                # vignette et échoue légitimement au contrôle de résolution des
                # médias iziLife.
                product['primary_image_url']=image if image.startswith(('http://','https://')) else f'https://img.qrfy.com/img/original/{image.lstrip("/")}'
            variants=[]
            if item.get('multiplePrices') and isinstance(item.get('prices'),list):
                for variant in item['prices']:
                    variant_name=str(variant.get('name') or '').strip()
                    variant_price=parse_price(variant.get('value'))
                    if variant_name and variant_price is not None: variants.append((variant_name,variant_price))
            volume_pattern=re.compile(r'^\s*(\d+(?:[.,]\d+)?)\s*(ml|cl|l)\s*$',re.IGNORECASE)
            volume_variants=bool(variants) and all(volume_pattern.match(label) for label,_ in variants)
            labels=[label.casefold() for label,_ in variants]
            count_variants=bool(variants) and all(re.search(r'\b(unité|unite|pièce|pieces|pièces|shooters?)\b',label,re.IGNORECASE) for label,_ in variants)
            group_name='Volume' if volume_variants else ('Quantité' if count_variants else None)
            # Automatisation stricte : seules les contenances et quantités
            # explicites sont créées. Une formule pouvant relever d'un Produit
            # composé nécessite une résolution métier manuelle.
            if len(variants) >= 2 and group_name:
                base=min(price for _,price in variants); product['price']=base
                product['option_groups']=[{
                    'name':group_name, 'selection_rule':1, 'min_choices':1, 'max_choices':1,
                    'values':[{
                        'name':normalize_quantity_label(variant_name),
                        'is_default':price == base,
                        'price_mode':'ABSOLUTE',
                        'amount':price,
                    } for variant_name,price in variants],
                }]
                if volume_variants:
                    base_label=min(variants,key=lambda pair: pair[1])[0]
                    measurement=volume_pattern.match(base_label)
                    product['quantity_type']=3
                    product['product_quantity_value']=float(measurement.group(1).replace(',','.'))
                    product['volume_unit_code']=measurement.group(2).lower()
                    product['quantity_description']=base_label
                elif count_variants:
                    product['quantity_type']=1
                    product['product_quantity_value']=1
                    product['quantity_description']=min(variants,key=lambda pair: pair[1])[0]
            elif len(variants) == 1:
                product['price']=variants[0][1]
                measurement=volume_pattern.match(variants[0][0])
                if measurement:
                    product['quantity_type']=3
                    product['product_quantity_value']=float(measurement.group(1).replace(',','.'))
                    product['volume_unit_code']=measurement.group(2).lower()
                    product['quantity_description']=variants[0][0]
                else:
                    product['quantity_description']=variants[0][0]
            elif variants:
                product['price']=min(price for _,price in variants)
            else:
                price=parse_price(item.get('price'))
                if price is not None: product['price']=price
            products.append(product)
    return products

def run(zone,env,file,dry_run,maximum,analyze_only=False,scope_refs=None,fixture_file=None,llm_provider=None,source_urls=None):
    load_env(); env=normalize_env(env); path=Path(file) if file else workbook_file(zone,env)
    if not path.exists(): raise SystemExit(f'Fichier introuvable : {path}. Lance --init.')
    wb=openpyxl.load_workbook(path); ws=wb['Sources'] if 'Sources' in wb.sheetnames else wb.active
    schema_changed=configure_sources_sheet(ws)
    if schema_changed and not analyze_only and not dry_run: configure_guide_sheet(wb); wb.save(path)
    headers=[str(c.value or '').strip().upper() for c in ws[1]]; status_col=headers.index('STATUS')+1; result_col=headers.index('LAST_RESULT')+1
    rows=[]
    for row in ws.iter_rows(min_row=2):
        values={headers[i]:row[i].value for i in range(min(len(headers),len(row))) if headers[i]}
        current_ref=str(values.get('SCOPE_REF') or '').strip()
        # Une analyse seule est un diagnostic explicite : elle doit pouvoir être
        # rejouée sur une ligne déjà marquée done sans modifier le classeur.
        if not current_ref or (not analyze_only and str(values.get('STATUS') or '').strip().lower() not in PROCESSABLE_STATUSES): continue
        if scope_refs and current_ref not in scope_refs: continue
        rows.append((row[0].row,values))
        if maximum>0 and len(rows)>=maximum: break
    if dry_run:
        for n,v in rows: print('[DRY RUN]',n,v.get('SCOPE_TYPE'),v.get('SCOPE_REF'),extract_urls(v),[str(p) for p in extract_images(v,path)])
        return
    fixtures={}
    if fixture_file:
        fixture_payload=json.loads(Path(fixture_file).read_text(encoding='utf-8'))
        fixtures=fixture_payload.get('sources',fixture_payload) if isinstance(fixture_payload,dict) else {}
    stats={'rows':0,'items':0,'skipped':0,'errors':0}; analysis_report=[]
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=False,args=['--disable-blink-features=AutomationControlled']); page=browser.new_page(locale='fr-FR')
        for row_number,values in rows:
            try:
                scope=str(values.get('SCOPE_TYPE') or '').strip().upper(); ref=str(values.get('SCOPE_REF') or '').strip(); urls=extract_urls(values); images=extract_images(values,path); llm_instructions=str(values.get('LLM_INSTRUCTIONS') or '').strip()
                if source_urls: urls=[url for url in urls if url in source_urls]
                if scope not in {'PLACE','SHOP'}: raise RuntimeError('SCOPE_TYPE doit être PLACE ou SHOP.')
                if not urls and not images: raise RuntimeError('Au moins une URL ou une image de menu est obligatoire.')
                # Validate the target before opening source pages, invoking OCR/LLM,
                # or attempting any upsert. This endpoint is a read-only lookup.
                target=izilife_post('/scraper/agentGetPlaceByStringId',{
                    'string_id':ref,'entity_type':scope
                },env,timeout=60)
                if not target.get('success'):
                    raise RuntimeError(f"Scope introuvable: {scope}/{ref} ({target.get('error',target)})")
                target_item=target.get('item') if isinstance(target.get('item'),dict) else {}
                print(f"Ligne {row_number} — scope vérifié : {scope}/{ref} — {target_item.get('name','')}",flush=True)
                row_items=row_skipped=0; details=[]; row_review=False; row_warnings=[]; row_image_errors=[]
                for image_path in images:
                    if analyze_only:
                        details.append({'image':image_path.name,'mode':'image_ocr','skipped':'analyse seule indisponible pour les images'})
                        continue
                    print(f'Ligne {row_number} — OCR image : {image_path.name}',flush=True)
                    try:
                        result=ingest_menu_image(image_path,scope,ref,env)
                        if not result.get('success'): raise RuntimeError(str(result.get('error',result)))
                        row_items+=int(result.get('inserted',0))+int(result.get('updated',0))
                        row_skipped+=int(result.get('skipped',0))
                        details.append({'image':image_path.name,'detected':int(result.get('products',0)),'mode':'image_ocr'})
                    except Exception as exc:
                        message=f'{image_path.name}: {exc}'
                        row_image_errors.append(message)
                        details.append({'image':image_path.name,'mode':'image_ocr','error':str(exc)})
                        print(f'AVERTISSEMENT OCR ligne {row_number} — {message}; poursuite avec les autres images.',flush=True)
                for url in urls:
                    if 'dokaa.fr' in (urlparse(url).hostname or '').lower():
                        products=scrape_dokaa_products(page,url)
                        print(f'Ligne {row_number} — Dokaa structuré : {len(products)} produit(s) actif(s)')
                        for item_index,item in enumerate(products,1):
                            print(f'  Produit {item_index}/{len(products)} : {item["name"]}',flush=True)
                            if analyze_only: continue
                            payload={'kind':'product','scope_type':scope,'scope_ref':ref,'source_url':url,'item':item}
                            result=izilife_post('/scraper/agentIngestSiteItem',{'payload':json.dumps(payload,ensure_ascii=False)},env,timeout=180)
                            if not result.get('success'): raise RuntimeError(f'{url}: {result.get("error",result)}')
                            if result.get('created') or result.get('updated'): row_items+=1
                            else: row_skipped+=1
                        details.append({'url':url,'detected':len(products),'mode':'dokaa_structured'})
                        continue
                    if url in fixtures:
                        extracted_items=list(fixtures[url] or [])
                        print(f'Ligne {row_number} — fixture locale : {len(extracted_items)} élément(s) — {url}',flush=True)
                        if analyze_only:
                            analysis_report.append({'scope_type':scope,'scope_ref':ref,'url':url,'chunk':1,'items':extracted_items})
                        else:
                            for extracted in extracted_items:
                                payload={'kind':extracted.get('kind'),'scope_type':scope,'scope_ref':ref,'source_url':url,'item':extracted.get('item',{})}
                                result=izilife_post('/scraper/agentIngestSiteItem',{'payload':json.dumps(payload,ensure_ascii=False)},env,timeout=180)
                                if not result.get('success'): raise RuntimeError(f'{url}: {result.get("error",result)}')
                                if result.get('created') or result.get('updated'): row_items+=1
                                else: row_skipped+=1
                        details.append({'url':url,'detected':len(extracted_items),'mode':'fixture_locale'})
                        continue
                    content=scrape_exact_url(page,url)
                    taxonomy_candidate=izilife_post('/scraper/agentAnalyzeSiteContent',{
                        'scope_type':scope,'scope_ref':ref,'source_url':url,'content':content[:18000],
                        'provider':llm_provider or '','analysis_mode':'taxonomy'
                    },env,timeout=240)
                    if not taxonomy_candidate.get('success'):
                        raise RuntimeError(f'{url} structure globale: {taxonomy_candidate.get("error",taxonomy_candidate)}')
                    taxonomy={'element_menus':taxonomy_candidate.get('element_menus',[])}
                    taxonomy_json=json.dumps(taxonomy,ensure_ascii=False,separators=(',',':'))
                    if analyze_only:
                        analysis_report.append({'scope_type':scope,'scope_ref':ref,'url':url,'analysis_mode':'taxonomy','taxonomy':taxonomy,'warnings':taxonomy_candidate.get('warnings',[])})
                    effective_instructions=(llm_instructions+'\nPLAN GLOBAL DE CATÉGORIES VALIDÉ POUR TOUS LES BLOCS: '+taxonomy_json).strip()
                    detected=0; chunks=content_chunks(content); url_items=[]; url_warnings=list(map(str,taxonomy_candidate.get('warnings',[]))); url_review=False
                    for chunk_index,chunk in enumerate(chunks,1):
                        print(f'Ligne {row_number} — analyse bloc {chunk_index}/{len(chunks)} : {url}',flush=True)
                        analyzed=None
                        for attempt in range(1,4):
                            candidate=izilife_post('/scraper/agentAnalyzeSiteContent',{'scope_type':scope,'scope_ref':ref,'source_url':url,'content':chunk,'provider':llm_provider or '','llm_instructions':effective_instructions},env,timeout=240)
                            if candidate.get('success'):
                                analyzed=candidate;break
                            if candidate.get('retryable') is False: break
                            if attempt<3:
                                print(f'  Nouvelle tentative {attempt+1}/3 pour le bloc {chunk_index}',flush=True)
                                time.sleep(attempt*2)
                        if analyzed is None: raise RuntimeError(f'{url} bloc {chunk_index}/{len(chunks)}: {candidate.get("error",candidate)}')
                        detected+=len(analyzed.get('items',[]))
                        extracted_items=analyzed.get('items',[])
                        url_items.extend(extracted_items)
                        url_warnings.extend(str(value) for value in analyzed.get('warnings',[]) if str(value).strip())
                        url_review=url_review or bool(analyzed.get('review_required'))
                        if analyze_only:
                            analysis_report.append({'scope_type':scope,'scope_ref':ref,'url':url,'chunk':chunk_index,'items':extracted_items,'review_required':bool(analyzed.get('review_required')),'warnings':analyzed.get('warnings',[])})
                    url_items,menu_warnings=consolidate_element_menus(url_items)
                    url_warnings.extend(menu_warnings)
                    url_items=dedupe_extracted_items(url_items)
                    service_items=[entry for entry in url_items if str(entry.get('kind') or '').lower()=='element_service']
                    if len(service_items)>1:
                        consolidation=izilife_post('/scraper/agentAnalyzeSiteContent',{
                            'scope_type':scope,'scope_ref':ref,'source_url':url,
                            'content':json.dumps(service_items,ensure_ascii=False),
                            'provider':llm_provider or '','analysis_mode':'consolidate_services'
                        },env,timeout=240)
                        if consolidation.get('success'):
                            consolidated_services=[entry for entry in consolidation.get('items',[]) if str(entry.get('kind') or '').lower()=='element_service']
                            if consolidated_services:
                                url_items=[entry for entry in url_items if str(entry.get('kind') or '').lower()!='element_service']+dedupe_extracted_items(consolidated_services)
                                url_warnings.extend(str(value) for value in consolidation.get('warnings',[]) if str(value).strip())
                                if analyze_only:
                                    analysis_report.append({'scope_type':scope,'scope_ref':ref,'url':url,'analysis_mode':'consolidate_services','items':consolidated_services,'warnings':consolidation.get('warnings',[])})
                        else:
                            url_warnings.append('Consolidation des Services indisponible; déduplication déterministe conservée.')
                    if url_review:
                        row_review=True; row_warnings.extend(url_warnings)
                        details.append({'url':url,'detected':detected,'chunks':len(chunks),'review_required':True,'warnings':url_warnings,'proposals':compact_review_items(url_items)})
                        print(f'REVIEW ligne {row_number} — {url}: décision métier attendue',flush=True)
                        continue
                    for item_index,extracted in enumerate(url_items,1):
                        print(f'  Élément {item_index}/{len(url_items)} : {extracted.get("kind")}',flush=True)
                        if analyze_only: continue
                        payload={'kind':extracted.get('kind'),'scope_type':scope,'scope_ref':ref,'source_url':url,'item':extracted.get('item',{})}
                        result=izilife_post('/scraper/agentIngestSiteItem',{'payload':json.dumps(payload,ensure_ascii=False)},env,timeout=180)
                        if not result.get('success'): raise RuntimeError(f'{url}: {result.get("error",result)}')
                        if result.get('created') or result.get('updated'): row_items+=1
                        else: row_skipped+=1
                    row_warnings.extend(url_warnings)
                    details.append({'url':url,'detected':detected,'unique':len(url_items),'chunks':len(chunks),'taxonomy':taxonomy,'warnings':url_warnings})
                if analyze_only:
                    stats['rows']+=1;stats['items']+=sum(int(d.get('detected',0)) for d in details);print(f'OK analyse seule ligne {row_number}')
                    continue
                final_status='review' if row_review else ('retry' if row_image_errors else 'done')
                result_payload={'items':row_items,'skipped':row_skipped,'review_required':row_review,'warnings':row_warnings,'image_errors':row_image_errors,'sources':details}
                result_text=json.dumps(result_payload,ensure_ascii=False)
                if len(result_text)>32000: result_text=json.dumps({'items':row_items,'skipped':row_skipped,'review_required':row_review,'warnings':row_warnings,'message':'Résultat détaillé tronqué; utiliser --analyze-only pour le rapport complet.'},ensure_ascii=False)
                ws.cell(row_number,status_col,final_status); ws.cell(row_number,result_col,result_text); wb.save(path)
                stats['rows']+=1;stats['items']+=row_items;stats['skipped']+=row_skipped;stats['errors']+=len(row_image_errors)
                if row_image_errors: print(f'RETRY ligne {row_number}: {row_items} objet(s), {len(row_image_errors)} image(s) en échec.')
                else: print(f'OK ligne {row_number}: {row_items} objet(s)')
            except Exception as exc:
                if not analyze_only:
                    ws.cell(row_number,status_col,'error');ws.cell(row_number,result_col,str(exc)[:32000]);wb.save(path)
                stats['errors']+=1;print(f'ERREUR ligne {row_number}: {exc}')
        browser.close()
    if analyze_only:
        report_path=root_dir(zone,env)/'analyse_seule_enrichisseur.json'
        report_path.write_text(json.dumps({'stats':stats,'results':analysis_report},ensure_ascii=False,indent=2),encoding='utf-8')
        print(f'Rapport analyse seule : {report_path}')
    print('Résultat:',json.dumps(stats,ensure_ascii=False))

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--zone',required=True);ap.add_argument('--env',choices=['local','staging','prod'],default='local');ap.add_argument('--file');ap.add_argument('--init',action='store_true');ap.add_argument('--dry-run',action='store_true');ap.add_argument('--analyze-only',action='store_true',help='Scrape et analyse les URL sans insérer ni modifier le classeur.');ap.add_argument('--fixture-file',help='JSON simulant les items LLM par URL, sans appeler le fournisseur IA.');ap.add_argument('--llm-provider',choices=['openai','claude'],help='Force le fournisseur pour l’analyse de contenu.');ap.add_argument('--scope-ref',action='append',dest='scope_refs',help='Limiter le traitement à une ou plusieurs références de scope.');ap.add_argument('--source-url',action='append',dest='source_urls',help='Limiter le diagnostic à une ou plusieurs URL exactes du classeur.');ap.add_argument('--max',type=int,default=100);a=ap.parse_args();init_sheet(a.zone,normalize_env(a.env)) if a.init else run(a.zone,a.env,a.file,a.dry_run,a.max,a.analyze_only,set(a.scope_refs or []),a.fixture_file,a.llm_provider,set(a.source_urls or []))

if __name__=='__main__':
    from agent_excel_logger import run_logged
    run_logged('enrichisseur_contenus','pilotage',main)
