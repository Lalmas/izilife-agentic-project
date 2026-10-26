from __future__ import annotations
import argparse,json,os,re,sys,unicodedata
from datetime import datetime
from pathlib import Path
from urllib.parse import quote_plus,urlparse
import openpyxl,requests

HERE=Path(__file__).resolve();sys.path.insert(0,str(HERE.parents[2]));sys.path.insert(0,str(HERE.parents[1]))
from core.paths import IZILIFE_ENVS,drive_workspace_root,normalize_env,normalize_zone
from agent_excel_logger import run_logged

HEADERS=["target_type","target_ref","target_title","site","product_url","status","notes","last_result"]
GYG_EXAMPLES=[
 ["EXPERIENCE","visiter-lille-en-2-cv","Visiter Lille en 2CV","getyourguide","https://www.getyourguide.com/lille-l643/visite-originale-de-lille-en-2cv-decapotable-t31803/?ranking_uuid=d6ab685f-24f3-4373-b698-3d6d35c5672e&referral_redirect=1","pending","", ""],
 ["EXPERIENCE","visite-guidee-du-vieux-lille","Visite guidée du Vieux-Lille","getyourguide","https://www.getyourguide.com/lille-l643/lille-visite-a-pied-de-2-heures-du-vieux-lille-en-francais-t50838/","pending","", ""],
]

def api_post(env_name,endpoint,data):
    token=(os.getenv("IZILIFE_AGENT_TOKEN") or os.getenv("AGENT_SECRET_TOKEN") or "").strip()
    if not token: raise RuntimeError("IZILIFE_AGENT_TOKEN ou AGENT_SECRET_TOKEN manquant.")
    env=IZILIFE_ENVS[normalize_env(env_name)]
    response=requests.post(env["base_url"].rstrip("/")+endpoint,data=data,headers={"X-Agent-Token":token},verify=env.get("verify_ssl",True),timeout=120)
    try: result=response.json()
    except Exception: result={"success":False,"error":response.text[:1000]}
    if response.status_code>=400: result.setdefault("success",False)
    return result

def workbook_path(env_name,zone,provider):
    return drive_workspace_root(env_name)/"izilife"/"affiliations"/normalize_zone(zone)/f"{provider}.xlsx"

def create_book(path,rows):
    path.parent.mkdir(parents=True,exist_ok=True)
    if path.exists(): print(f"Existe déjà : {path}");return
    wb=openpyxl.Workbook();ws=wb.active;ws.title="Affiliations";ws.append(HEADERS)
    for row in rows: ws.append(row)
    ws.freeze_panes="A2";ws.auto_filter.ref=f"A1:H{max(2,len(rows)+1)}";wb.save(path);print(f"Créé : {path}")

def clean(value):
    value=unicodedata.normalize("NFKD",str(value or "")).encode("ascii","ignore").decode().lower()
    return " ".join(re.findall(r"[a-z0-9]+",value))

def ticketmaster_candidates(page,item):
    candidates=[]
    title=item["title"];venue=item.get("venue_name","");city=item.get("city_name","")
    search_urls=[
        "https://www.ticketmaster.fr/fr/resultat?ipSearch="+quote_plus(title),
        "https://www.google.com/search?q="+quote_plus(f'site:ticketmaster.fr/fr/manifestation "{title}" "{city}"'),
        "https://www.google.com/search?q="+quote_plus(f'site:ticketmaster.fr/fr/manifestation "{title}" {venue} {city}'),
    ]
    for search_url in search_urls:
        page.goto(search_url,wait_until="domcontentloaded",timeout=60000)
        if "ticketmaster.fr/fr/resultat" in search_url: page.wait_for_timeout(2500)
        for anchor in page.locator("a").all():
            href=anchor.get_attribute("href") or "";text=anchor.inner_text().strip()
            if href.startswith("/url?q="): href=href[7:].split("&",1)[0]
            host=(urlparse(href).hostname or "").lower()
            if not(host=="ticketmaster.fr" or host.endswith(".ticketmaster.fr")): continue
            haystack=clean(text+" "+href)
            title_tokens=[x for x in clean(title).split() if len(x)>2]
            venue_tokens=[x for x in clean(venue).split() if len(x)>3]
            city_tokens=[x for x in clean(city).split() if len(x)>2]
            title_score=sum(x in haystack for x in title_tokens)/max(1,len(title_tokens))
            venue_score=sum(x in haystack for x in venue_tokens)/max(1,len(venue_tokens)) if venue_tokens else 0
            city_score=sum(x in haystack for x in city_tokens)/max(1,len(city_tokens)) if city_tokens else 0
            location_score=max(venue_score,city_score)
            candidates.append({"url":href.split("?",1)[0],"label":text,"score":round(title_score*.7+location_score*.3,3),"title_score":round(title_score,3),"location_score":round(location_score,3)})
    unique={}
    for item in candidates:
        if item["url"] not in unique or item["score"]>unique[item["url"]]["score"]: unique[item["url"]]=item
    return sorted(unique.values(),key=lambda row:row["score"],reverse=True)

def run_ticketmaster(args):
    response=api_post(args.env,"/scraper/agentAffiliateCandidates",{"site":"ticketmaster","limit":args.max})
    if not response.get("success"): raise RuntimeError(response.get("error","Récupération des Event impossible."))
    items=response.get("items",[]);output=[]
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        browser=pw.chromium.launch(headless=not args.visible);page=browser.new_page(locale="fr-FR",timezone_id="Europe/Paris")
        for event in items:
            matches=ticketmaster_candidates(page,event);best=matches[0] if matches else None;status="not_found";result={"success":False,"reason":"Aucun résultat Ticketmaster"}
            if best and best["score"]>=args.min_score and best["title_score"]>=.75 and best["location_score"]>=.5:
                payload={"target_type":event["target_type"],"target_id":event["target_id"],"site":"ticketmaster","url":best["url"]}
                result={"success":True,"dry_run":True,"payload":payload} if args.dry_run else api_post(args.env,"/scraper/agentUpsertAffiliateLink",{"payload":json.dumps(payload,ensure_ascii=False)})
                status="done" if result.get("success") else "error"
            elif best: status="review";result={"success":False,"reason":"Score insuffisant","candidate":best}
            output.append([event["target_type"],event.get("target_ref") or event["target_id"],event["title"],"ticketmaster",best["url"] if best else "",status,f'{event.get("event_start_date","")} · {event.get("venue_name","")}',json.dumps(result,ensure_ascii=False)[:2000]])
            print(status.upper(),event["title"],best or "")
        browser.close()
    path=workbook_path(args.env,args.zone,"ticketmaster").with_name(f"ticketmaster_{datetime.now():%Y%m%d_%H%M%S}.xlsx");create_book(path,output)
    print(f"Traités: {len(items)}");print(f"Créés: {sum(1 for row in output if row[5]=='done')}");print(f"Erreurs: {sum(1 for row in output if row[5]=='error')}")

def run_getyourguide(args):
    path=workbook_path(args.env,args.zone,"getyourguide")
    if args.init: create_book(path,GYG_EXAMPLES);return
    if not path.exists(): create_book(path,GYG_EXAMPLES)
    wb=openpyxl.load_workbook(path);ws=wb["Affiliations"];headers=[cell.value for cell in ws[1]];processed=created=errors=0
    for number in range(2,ws.max_row+1):
        row={headers[index]:ws.cell(number,index+1).value for index in range(len(headers))}
        if str(row.get("status") or "pending").lower() in {"done","skip"}: continue
        payload={"target_type":row["target_type"],"target_ref":row["target_ref"],"site":"getyourguide","url":row["product_url"]}
        result={"success":True,"dry_run":True,"payload":payload} if args.dry_run else api_post(args.env,"/scraper/agentUpsertAffiliateLink",{"payload":json.dumps(payload,ensure_ascii=False)})
        processed+=1
        if result.get("success"): created+=int(bool(result.get("created")));ws.cell(number,headers.index("status")+1,"done")
        else: errors+=1;ws.cell(number,headers.index("status")+1,"error")
        ws.cell(number,headers.index("last_result")+1,json.dumps(result,ensure_ascii=False)[:2000]);print("OK" if result.get("success") else "ERREUR",row["target_ref"],result)
    if not args.dry_run: wb.save(path)
    print(f"Traités: {processed}");print(f"Créés: {created}");print(f"Erreurs: {errors}")

def main():
    p=argparse.ArgumentParser(description="Recherche et insertion des liens affiliés iziLife.");p.add_argument("--provider",choices=["ticketmaster","getyourguide"],required=True);p.add_argument("--zone",required=True);p.add_argument("--env",choices=["local","staging","prod"],default="local");p.add_argument("--init",action="store_true");p.add_argument("--dry-run",action="store_true");p.add_argument("--visible",action="store_true");p.add_argument("--max",type=int,default=100);p.add_argument("--min-score",type=float,default=.78);args=p.parse_args()
    run_ticketmaster(args) if args.provider=="ticketmaster" else run_getyourguide(args)

if __name__=="__main__":
    provider=next((sys.argv[i+1] for i,arg in enumerate(sys.argv[:-1]) if arg=="--provider"),"affiliate");run_logged(f"affiliate_{provider}","affiliations",main)
