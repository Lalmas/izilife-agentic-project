"""Agent iziLife : synchronise les éditions sportives sans ID technique ni accès serveur."""
import argparse, csv, json, os, re, sys, time
from datetime import date, datetime, timedelta
from pathlib import Path
import requests
import urllib3
from dotenv import load_dotenv

WORKSPACE_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(WORKSPACE_ROOT / "scripts"))
from core.paths import ENV_GLOBAL, ENV_IZILIFE, IZILIFE_ENVS
try:
    from playwright.sync_api import sync_playwright
except ImportError:
    print("Installez : pip install -r requirements-competition.txt")
    sys.exit(1)

for env_file in (ENV_GLOBAL, ENV_IZILIFE):
    if env_file.exists(): load_dotenv(env_file, override=False)
ENVS = IZILIFE_ENVS
TOKEN = os.environ.get("IZILIFE_AGENT_TOKEN", "").strip()
if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8", errors="replace")
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
DEFAULT_FILE = Path(__file__).resolve().parent / "competition-planning.xlsx"

def log(message): print(f"[{datetime.now().strftime('%H:%M:%S')}] {message}")
def truthy(value, default=True):
    return default if value is None or str(value).strip() == "" else str(value).strip().lower() in {"1","oui","yes","true","active"}
def iso(value):
    if value in (None, ""): return ""
    return value.strftime("%Y-%m-%d") if isinstance(value, (date, datetime)) else str(value).strip()[:10]
def normalize(row):
    row = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
    return {
        "active": truthy(row.get("active")),
        "competition_string_id": str(row.get("competition_string_id") or row.get("competition") or "").strip(),
        "season_label": str(row.get("season_label") or "").strip(),
        "edition_name": str(row.get("edition_name") or "").strip(),
        "starts_at": iso(row.get("starts_at")), "ends_at": iso(row.get("ends_at")),
        "schedule_source_url": str(row.get("schedule_source_url") or row.get("calendar_url") or "").strip(),
        "import_mode": str(row.get("import_mode") or "AUTO").strip().upper(),
    }

def init_workbook(filename):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(filename)
    if path.exists():
        log(f"Existe déjà : {path}")
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Éditions"
    headers = ["active","competition_string_id","season_label","edition_name","starts_at","ends_at","schedule_source_url","import_mode"]
    sheet.append(headers)
    samples = [
        ["Oui","ligue-1","","","","","https://ligue1.com/fr/calendar/ligue1","AUTO"],
        ["Oui","ligue-2","","","","","https://ligue1.com/fr/calendar/ligue2","AUTO"],
        ["Oui","ligue-des-champions-de-football","","","","","","SKIP"],
        ["Oui","starligue","","","","","","SKIP"],
        ["Oui","roland-garros","","","","","","SKIP"],
    ]
    for row in samples: sheet.append(row)
    widths = [11,36,20,30,15,15,70,15]
    for index, width in enumerate(widths, 1):
        cell = sheet.cell(1, index); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173F5F"); cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[cell.column_letter].width = width
    for row in sheet.iter_rows(min_row=2, max_row=200, min_col=3, max_col=7):
        for cell in row: cell.fill = PatternFill("solid", fgColor="FFF7D6")
    active_validation = DataValidation(type="list", formula1='"Oui,Non"')
    mode_validation = DataValidation(type="list", formula1='"AUTO,FORCE,SKIP"')
    sheet.add_data_validation(active_validation); active_validation.add("A2:A200")
    sheet.add_data_validation(mode_validation); mode_validation.add("H2:H200")
    sheet.freeze_panes = "C2"; sheet.auto_filter.ref = "A1:H200"
    help_sheet = workbook.create_sheet("Aide")
    help_sheet.append(["Commande", "python fetch_competition_schedule.py --env prod"])
    help_sheet.append(["AUTO", "Crée ou met à jour seulement si nécessaire."])
    help_sheet.append(["FORCE", "Relit le calendrier même si l'édition est déjà à jour."])
    help_sheet.append(["SKIP", "Prépare l'édition sans importer les matchs."])
    help_sheet.column_dimensions["A"].width = 20; help_sheet.column_dimensions["B"].width = 80
    workbook.save(path); log(f"Créé : {path}")

def load_rows(filename):
    path = Path(filename)
    if not path.exists(): raise RuntimeError(f"Fichier introuvable : {path}")
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [normalize(row) for row in csv.DictReader(handle)]
    if path.suffix.lower() not in {".xlsx", ".xlsm"}: raise RuntimeError("Format attendu : xlsx, xlsm ou csv")
    try: from openpyxl import load_workbook
    except ImportError as exc: raise RuntimeError("openpyxl manque : pip install -r requirements.txt") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Éditions"] if "Éditions" in workbook.sheetnames else workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values: return []
    headers = [str(v or "").strip() for v in values[0]]
    return [normalize(dict(zip(headers, row))) for row in values[1:] if any(v not in (None, "") for v in row)]

def post(url, data, env, timeout=120):
    response = requests.post(url, data=data, headers={"X-Agent-Token": TOKEN}, verify=env["verify_ssl"], timeout=timeout)
    try: result = response.json()
    except ValueError: result = {"success": False, "error": response.text[:500]}
    if response.status_code >= 400:
        result.update(success=False, error=f"HTTP {response.status_code} — {result.get('error','erreur API')}")
    return result

def prepare(row, env, dry_run):
    data = {key: row[key] for key in ("competition_string_id","season_label","edition_name","starts_at","ends_at","schedule_source_url")}
    data["dry_run"] = "1" if dry_run else "0"
    return post(f"{env['base_url']}/scraper/prepareCompetitionEditionFromAgent", data, env, 30)

def load_internal_plan(env):
    result = post(f"{env['base_url']}/scraper/competitionSyncPlanFromAgent", {}, env, 30)
    if not result.get("success"): raise RuntimeError(result.get("error", "plan interne indisponible"))
    return [normalize(row) for row in result.get("competitions", [])]

def fetch_page(page, url):
    log(f"Lecture → {url}")
    try: page.goto(url, wait_until="networkidle", timeout=45000)
    except Exception: pass
    try:
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)"); time.sleep(1)
        text=page.locator("body").inner_text(timeout=15000)
    except Exception: text=""
    if len(text.strip())>=200:return text
    try:
        response=requests.get(url,headers={"User-Agent":"Mozilla/5.0 (compatible; iziLife schedule agent)"},timeout=30)
        response.raise_for_status()
        raw=re.sub(r"<script[\s\S]*?</script>|<style[\s\S]*?</style>"," ",response.text,flags=re.I)
        return re.sub(r"\s+"," ",re.sub(r"<[^>]+>"," ",raw)).strip()
    except Exception:return text

def parse_ligue(text, edition):
    if "ligue1.com/fr/calendar/" not in str(edition.get("source_url", "")): return []
    lines = [v.strip() for v in text.splitlines() if v.strip()]
    months={"jan":1,"fév":2,"fev":2,"mar":3,"avr":4,"mai":5,"juin":6,"juil":7,"août":8,"aout":8,"sept":9,"oct":10,"nov":11,"déc":12,"dec":12}
    start_year=int(str(edition.get("starts_at") or datetime.now().year)[:4]); start_month=int(str(edition.get("starts_at") or "0000-08")[5:7])
    current=None; phase="Journée"; matches=[]
    for index,line in enumerate(lines):
        if re.match(r"^Journée\s+\d+$",line,re.I): phase=line
        if line.lower()=="aujourd'hui": current=datetime.now().strftime("%Y-%m-%d 00:00:00")
        elif line.lower()=="hier": current=(datetime.now()-timedelta(days=1)).strftime("%Y-%m-%d 00:00:00")
        dm=re.search(r"(\d{1,2})\s+([A-Za-zÀ-ÿ]+)",line)
        if dm:
            key=dm.group(2).lower().rstrip(".")[:4]; month=next((v for k,v in months.items() if key.startswith(k) or k.startswith(key)),None)
            if month: current=f"{start_year if month>=start_month else start_year+1:04d}-{month:02d}-{int(dm.group(1)):02d} 00:00:00"
        if re.fullmatch(r"-\s+-|\d+\s*-\s*\d+|\d{1,2}H(?:\d{2})?",line,re.I) and current:
            before=[v for v in lines[max(0,index-3):index] if not re.fullmatch(r"\(\d+\)",v)]; after=[v for v in lines[index+1:index+4] if not re.fullmatch(r"\(\d+\)",v)]
            if before and after:
                home,away=before[-1],after[0]
                if home not in {"Classement","Matchs"} and away not in {"Classement","Voir le classement complet"}:
                    starts=current; tm=re.fullmatch(r"(\d{1,2})H(\d{2})?",line,re.I)
                    if tm: starts=starts[:11]+f"{int(tm.group(1)):02d}:{int(tm.group(2) or 0):02d}:00"
                    matches.append({"match_title":f"{home} - {away}","participant_1_name":home,"participant_2_name":away,"starts_at":starts,"phase_label":phase,"is_featured_home":0})
    return matches

def import_schedule(edition, env):
    if not edition.get("source_url"): return {"success":False,"error":"aucune URL sur l'édition ni sur la compétition"}
    source_url=edition["source_url"]
    if "ligue1.com/fr/calendar/" not in source_url:
        with sync_playwright() as playwright:
            browser=playwright.chromium.launch(headless=True)
            page=browser.new_page(locale="fr-FR", timezone_id="Europe/Paris")
            text=fetch_page(page,source_url)
            browser.close()
        if len(text.strip())<200:return {"success":False,"error":"contenu officiel vide ou inaccessible"}
        log("Source sans connecteur déterministe : extraction contrôlée côté iziLife")
        return post(f"{env['base_url']}/scraper/importCompetitionScheduleFromAgent/{edition['edition_id']}",{"html":text},env)
    from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
    parts=urlsplit(source_url); query=dict(parse_qsl(parts.query)); all_matches=[]; texts=[]
    with sync_playwright() as playwright:
        browser=playwright.chromium.launch(headless=True); page=browser.new_page(locale="fr-FR", timezone_id="Europe/Paris")
        for gameweek in range(1,35):
            query["gameweek"]=str(gameweek)
            url=urlunsplit((parts.scheme,parts.netloc,parts.path,urlencode(query),parts.fragment))
            text=fetch_page(page,url); texts.append(text)
            page_edition=dict(edition,source_url=url)
            found=parse_ligue(text,page_edition)
            if not found: log(f"Journée {gameweek}: aucun match reconnu")
            all_matches.extend(found)
        browser.close()
    unique={}
    for match in all_matches:
        key=(match.get("starts_at"),match.get("participant_1_name"),match.get("participant_2_name"))
        unique[key]=match
    matches=list(unique.values()); html="\n\n".join(texts)
    if not matches: return {"success":False,"error":"aucun match Ligue 1/Ligue 2 reconnu sur les 34 journées"}
    log(f"{len(matches)} match(s) uniques reconnus sur 34 journées")
    return post(f"{env['base_url']}/scraper/importCompetitionScheduleFromAgent/{edition['edition_id']}",{"html":html,"matches_json":json.dumps(matches,ensure_ascii=False)},env)

def main():
    parser=argparse.ArgumentParser(description="Upsert annuel des éditions et calendriers sportifs")
    source=parser.add_mutually_exclusive_group(); source.add_argument("--competition",help="string_id, ex. ligue-1"); source.add_argument("--file",help="Excel/CSV, une édition par ligne"); source.add_argument("--all",action="store_true",help="toutes les compétitions automatiques configurées en base")
    parser.add_argument("--season",default=""); parser.add_argument("--url",default=""); parser.add_argument("--env",choices=ENVS,default="local")
    parser.add_argument("--init",action="store_true",help="crée le fichier Excel modèle puis s'arrête")
    parser.add_argument("--force",action="store_true"); parser.add_argument("--dry-run",action="store_true"); args=parser.parse_args()
    file_path = Path(args.file) if args.file else DEFAULT_FILE
    if args.init:
        init_workbook(file_path); return 0
    if not TOKEN: log("IZILIFE_AGENT_TOKEN non défini."); return 2
    env=ENVS[args.env]
    rows=(load_internal_plan(env) if args.all else ([normalize({"competition_string_id":args.competition,"season_label":args.season,"schedule_source_url":args.url,"import_mode":"FORCE" if args.force else "AUTO"})] if args.competition else load_rows(file_path)))
    failures=0
    for row in rows:
        sid=row["competition_string_id"]
        if not row["active"] or not sid: continue
        log(f"{sid} — préparation"); edition=prepare(row,env,args.dry_run)
        if not edition.get("success"): failures+=1; log(f"ERREUR : {edition.get('error','inconnue')}"); continue
        log(f"{edition['action']} — {edition['name']} — source={edition.get('source_mode','?')} — cadence={edition.get('sync_cadence','?')}"); mode=row["import_mode"]
        should=not args.dry_run and mode!="SKIP" and (args.force or mode=="FORCE" or edition.get("needs_import"))
        if not should: log("rien à faire"); continue
        source_mode=str(edition.get("source_mode") or "").upper()
        try:
            if source_mode=="API_SPORTS":
                result=post(f"{env['base_url']}/scraper/importCompetitionScheduleApiFromAgent/{edition['edition_id']}",{},env)
                if not result.get("success") and edition.get("source_url"):
                    log(f"API-Sports indisponible — fallback Playwright ({result.get('error','erreur API')})")
                    result=import_schedule(edition,env)
            elif source_mode=="PLAYWRIGHT":
                result=import_schedule(edition,env)
            else:
                result={"skipped_source":True,"reason":"source absente ou non configurée"}
        except Exception as exc:
            result={"success":False,"error":f"{type(exc).__name__}: {exc}"}
        if result.get("skipped_source"): log(f"IGNORÉ — {result.get('reason')}")
        elif result.get("success"):
            log(f"OK — {result.get('inserted',0)} créé(s), {result.get('skipped',0)} mis à jour/ignoré(s)")
            enriched=post(f"{env['base_url']}/scraper/enrichCompetitionTeamsFromAgent",{"edition_id":edition["edition_id"]},env,60)
            if enriched.get("success") and not enriched.get("skipped"): log(f"LOGOS — {enriched.get('updated',0)} équipe(s) enrichie(s)")
            elif enriched.get("skipped"): log(f"LOGOS ignorés — {enriched.get('reason','non configuré')}")
            else: log(f"LOGOS indisponibles — {enriched.get('error','erreur')}")
        else: failures+=1; log(f"ERREUR import : {result.get('error','inconnue')}")
    return 1 if failures else 0
if __name__=="__main__": sys.exit(main())
