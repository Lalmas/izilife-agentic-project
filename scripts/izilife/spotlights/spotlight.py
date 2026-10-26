from __future__ import annotations
import argparse,sys
from pathlib import Path
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"offers"))
from offer_agent_common import create_workbook,drive_workspace_root,normalize_zone,process

FILE,SHEET="spotlights.xlsx","Spotlights"
CONTEXT_FIELDS=("days_of_week","visit_moments","start_month","end_month","visit_weathers")
COLUMNS=[("scope_type","scope"),("scope_ref","scope"),("city_ref","scope"),("editorial_title","offer"),("editorial_description","offer"),("theme","offer"),("target_spotlight","access"),("target_perfect_for","access"),("target_today","access"),("target_next_days","access"),("rail_editorial_reason","access"),("rail_configuration_json","access"),("reasons_json","offer")]
# Five explicit GOOD_FOR slots keep each reason independently configurable without
# requiring users to author JSON in Excel. `reasons_json` remains available for WHY
# and other legacy reason fields; slot entries are merged into that array.
for _slot in range(1,6):
    COLUMNS.extend([
        (f"good_for_{_slot}_type","offer"),(f"good_for_{_slot}_ref","offer"),
        (f"good_for_{_slot}_weight","offer"),(f"good_for_{_slot}_editorial_title","offer"),
        (f"good_for_{_slot}_editorial_description","offer"),
        *[(f"good_for_{_slot}_{field}","access") for field in CONTEXT_FIELDS],
    ])
COLUMNS.extend([("start_date","access"),("end_date","access"),("priority","bpr"),("base_weight","bpr"),("publication_status","bpr"),("status","result"),("notes","result"),("last_result","result")])
COLORS={"scope":"4472C4","offer":"70AD47","access":"FFC000","bpr":"A64D79","result":"7F8C8D"}

def ensure_columns(path):
    if not path.exists(): return
    import openpyxl
    wb=openpyxl.load_workbook(path); ws=wb[SHEET]
    headers=[str(c.value or "").strip() for c in ws[1]]
    added=False
    for name,group in COLUMNS:
        if name in headers: continue
        col=ws.max_column+1; cell=ws.cell(1,col,name)
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor=COLORS[group]); cell.alignment=Alignment(wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width=max(14,min(36,len(name)+3))
        if name.endswith("_type") and name.startswith("good_for_"):
            validation=DataValidation(type="list",formula1='"HOBBY,LITTLE_ACTIVITY_OCCASION"',allow_blank=True)
            ws.add_data_validation(validation); validation.add(f"{get_column_letter(col)}2:{get_column_letter(col)}500")
        headers.append(name); added=True
    if added:
        ws.auto_filter.ref=f"A1:{get_column_letter(ws.max_column)}{max(500,ws.max_row)}"
        wb.save(path)

def as_bool(value):
    return int(str(value or "").strip().lower() in {"1","true","yes","oui","y","x"})

def _list(value, field):
    if value in (None, ""):
        return []
    values=[part.strip() for part in str(value).replace(";",",").split(",") if part.strip()]
    if field == "days_of_week":
        result=[]
        for part in values:
            day=int(part)
            if day < 1 or day > 7: raise ValueError("jours de semaine: valeurs attendues de 1 (lundi) à 7 (dimanche).")
            if day not in result: result.append(day)
        return result
    return list(dict.fromkeys(values))

def _context(row, prefix=""):
    context={}
    days=_list(row.get(prefix+"days_of_week"),"days_of_week")
    moments=_list(row.get(prefix+"visit_moments"),"visit_moments")
    weathers=_list(row.get(prefix+"visit_weathers"),"visit_weathers")
    if days: context["days_of_week"]=days
    if moments: context["visit_moments"]=moments
    start=row.get(prefix+"start_month"); end=row.get(prefix+"end_month")
    if (start in (None, "")) != (end in (None, "")):
        raise ValueError("mois de début et mois de fin doivent être renseignés ensemble.")
    if start not in (None, ""):
        start=int(start); end=int(end)
        if not 1 <= start <= 12 or not 1 <= end <= 12: raise ValueError("les mois doivent être compris entre 1 et 12.")
        context["start_month"]=start; context["end_month"]=end
    if weathers: context["visit_weathers"]=weathers
    return context or None

def _reasons(row):
    raw=row.get("reasons_json")
    if raw in (None, ""):
        reasons=[]
    elif isinstance(raw,list):
        reasons=raw
    else:
        import json
        reasons=json.loads(str(raw))
    if not isinstance(reasons,list): raise ValueError("reasons_json doit contenir une liste si renseigné.")
    for reason in reasons:
        if not isinstance(reason,dict): continue
        if str(reason.get("role") or "GOOD_FOR").upper()=="WHY":
            reason.pop("context_json",None)
        else:
            # Preserve an explicit legacy context_json; otherwise contexts are
            # supplied per reason via the corresponding Excel slot below.
            reason.setdefault("context_json",None)
    for slot in range(1,6):
        ref=row.get(f"good_for_{slot}_ref")
        kind=row.get(f"good_for_{slot}_type")
        context=_context(row,f"good_for_{slot}_")
        if ref in (None, ""):
            if kind not in (None, "") or context is not None:
                raise ValueError(f"GOOD_FOR {slot}: type/ref obligatoire dès qu'un contexte est renseigné.")
            continue
        if kind in (None, ""):
            raise ValueError(f"GOOD_FOR {slot}: renseigner le type (HOBBY ou LITTLE_ACTIVITY_OCCASION).")
        reason={"type":str(kind).strip().upper(),"ref":str(ref).strip(),"role":"GOOD_FOR"}
        if row.get(f"good_for_{slot}_weight") not in (None, ""): reason["weight"]=int(row[f"good_for_{slot}_weight"])
        if row.get(f"good_for_{slot}_editorial_title") not in (None, ""): reason["editorial_title"]=row[f"good_for_{slot}_editorial_title"]
        if row.get(f"good_for_{slot}_editorial_description") not in (None, ""): reason["editorial_description"]=row[f"good_for_{slot}_editorial_description"]
        if context is not None: reason["context_json"]=context
        reasons.append(reason)
    return reasons or None

def payload(r):
    destinations = [as_bool(r.get(name)) for name in ("target_spotlight","target_perfect_for","target_today","target_next_days")]
    if not any(destinations): raise ValueError("Au moins une destination doit être activée.")
    return {"target_type":str(r.get("scope_type") or "").upper(),"target_string_id":r.get("scope_ref"),"city_string_id":r.get("city_ref"),"editorial_title":r.get("editorial_title"),"editorial_description":r.get("editorial_description"),"theme":r.get("theme"),"spotlight":destinations[0],"perfect_for":destinations[1],"rail_today":destinations[2],"rail_future":destinations[3],"rail_reason":r.get("rail_editorial_reason"),"rail_configuration_json":r.get("rail_configuration_json"),"reasons_json":_reasons(r),"start_date":r.get("start_date"),"end_date":r.get("end_date"),"priority":r.get("priority") or 0,"base_weight":r.get("base_weight") or 1,"status":r.get("publication_status") or "DRAFT"}

def main():
    ap=argparse.ArgumentParser(description="Spotlight classique/éditorial, qualifications Parfait pour et curation conditionnelle des rails Aujourd'hui/Prochains jours."); ap.add_argument("--zone",required=True); ap.add_argument("--env",choices=["local","staging","prod"],default="staging"); ap.add_argument("--init",action="store_true"); ap.add_argument("--dry-run",action="store_true"); ap.add_argument("--max",type=int,default=100); a=ap.parse_args()
    path=drive_workspace_root(a.env)/"izilife"/"spotlights"/normalize_zone(a.zone)/FILE
    if a.init:
        ex={"scope_type":"EVENT","scope_ref":"evenement-exemple-lille","city_ref":"lille","editorial_title":"Un événement à découvrir","editorial_description":"Texte éditorial facultatif selon la destination choisie.","theme":"sorties","target_spotlight":"oui","target_perfect_for":"non","target_today":"non","target_next_days":"non","reasons_json":"","priority":10,"base_weight":1,"publication_status":"DRAFT","status":"skip","notes":"Pour chaque GOOD_FOR, renseigner un slot (type + ref) et ses colonnes de contexte lisibles. jours: 1=lundi … 7=dimanche; listes séparées par virgules/points-virgules; moments/météos = string_id. Les colonnes de contexte ne modifient jamais la configuration globale du Spotlight classique. WHY reste dans reasons_json et n'a pas de contexte.","last_result":""}
        yes_no=["oui","non"]
        validations={"scope_type":["PLACE","SHOP","EVENT","EVENT_SERIE","EXPERIENCE"],"target_spotlight":yes_no,"target_perfect_for":yes_no,"target_today":yes_no,"target_next_days":yes_no,"publication_status":["DRAFT","PUBLISHED","ARCHIVED"],"status":["pending","done","error","skip"]}
        validations.update({f"good_for_{slot}_type":["HOBBY","LITTLE_ACTIVITY_OCCASION"] for slot in range(1,6)})
        create_workbook(path,SHEET,COLUMNS,ex,validations); ensure_columns(path); return
    ensure_columns(path)
    process(path,SHEET,a.env,"/scraper/agentUpsertSpotlight",payload,a.dry_run,a.max)

if __name__=="__main__":
    sys.path.insert(0,str(Path(__file__).resolve().parents[1])); from agent_excel_logger import run_logged; run_logged("spotlights","spotlights",main)
