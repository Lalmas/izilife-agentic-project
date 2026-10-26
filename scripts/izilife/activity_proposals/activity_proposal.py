from __future__ import annotations
import argparse,sys
from pathlib import Path
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"offers"))
from offer_agent_common import create_workbook,drive_workspace_root,normalize_zone,process
FILE,SHEET="activity_proposals.xlsx","ActivityProposals"
CONTEXT_FIELDS=("days_of_week","visit_moments","start_month","end_month","visit_weathers")
COLUMNS=[("scope_type","scope"),("scope_ref","scope"),("occasion_ref","offer"),("action_1","offer"),("action_2","offer"),("action_3","offer"),("title_template","offer"),("contextual_description","offer")]+[(name,"access") for name in CONTEXT_FIELDS]+[("special_activity_show","access"),("is_active","bpr"),("status","result"),("notes","result"),("last_result","result")]
COLORS={"scope":"4472C4","offer":"70AD47","access":"FFC000","bpr":"A64D79","result":"7F8C8D"}

def ensure_columns(path):
    if not path.exists(): return
    import openpyxl
    wb=openpyxl.load_workbook(path); ws=wb[SHEET]
    headers=[str(c.value or "").strip() for c in ws[1]]; added=False
    for name,group in COLUMNS:
        if name in headers: continue
        col=ws.max_column+1; cell=ws.cell(1,col,name)
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor=COLORS[group]); cell.alignment=Alignment(wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width=max(14,min(36,len(name)+3))
        headers.append(name); added=True
    if added:
        ws.auto_filter.ref=f"A1:{get_column_letter(ws.max_column)}{max(500,ws.max_row)}"
        wb.save(path)

def _list(value, numeric=False):
    if value in (None, ""): return []
    values=[part.strip() for part in str(value).replace(";",",").split(",") if part.strip()]
    if not numeric: return list(dict.fromkeys(values))
    result=[]
    for part in values:
        day=int(part)
        if day < 1 or day > 7: raise ValueError("jours de semaine: valeurs attendues de 1 (lundi) à 7 (dimanche).")
        if day not in result: result.append(day)
    return result

def _context(r):
    context={}
    days=_list(r.get("days_of_week"),True); moments=_list(r.get("visit_moments")); weathers=_list(r.get("visit_weathers"))
    if days: context["days_of_week"]=days
    if moments: context["visit_moments"]=moments
    if weathers: context["visit_weathers"]=weathers
    start=r.get("start_month"); end=r.get("end_month")
    if (start in (None, "")) != (end in (None, "")): raise ValueError("mois de début et mois de fin doivent être renseignés ensemble.")
    if start not in (None, ""):
        start=int(start); end=int(end)
        if not 1 <= start <= 12 or not 1 <= end <= 12: raise ValueError("les mois doivent être compris entre 1 et 12.")
        context["start_month"]=start; context["end_month"]=end
    return context or None

def payload(r):
    data={k:r.get(k) for k in ["scope_type","scope_ref","occasion_ref","action_1","action_2","action_3","title_template","contextual_description","special_activity_show","is_active"]}
    data["context_json"]=_context(r)
    return data
def main():
    ap=argparse.ArgumentParser(description="Propositions d'activité contextualisées par cible.");ap.add_argument("--zone",required=True);ap.add_argument("--env",choices=["local","staging","prod"],default="local");ap.add_argument("--init",action="store_true");ap.add_argument("--dry-run",action="store_true");ap.add_argument("--max",type=int,default=100);a=ap.parse_args()
    path=drive_workspace_root(a.env)/"izilife"/"activity-proposals"/normalize_zone(a.zone)/FILE
    if a.init:
        ex={"scope_type":"PLACE","scope_ref":"parc-de-la-dreve","occasion_ref":"date","action_1":"se-poser","action_2":"pique-niquer","title_template":"Se poser et pique-niquer au parc","contextual_description":"Une idée simple pour profiter du lieu.","is_active":"oui","status":"skip","notes":"jours: 1=lundi … 7=dimanche; listes séparées par virgules/points-virgules; moments/météos = string_id. Laisser vide pour context_json=null."}
        create_workbook(path,SHEET,COLUMNS,ex,{"scope_type":["PLACE","SHOP","EVENT","EVENT_SERIE","EXPERIENCE","ANNUAL_CELEBRATION"],"is_active":["oui","non"],"status":["pending","done","error","skip"]});ensure_columns(path);return
    ensure_columns(path)
    process(path,SHEET,a.env,"/scraper/agentUpsertActivityProposal",payload,a.dry_run,a.max)
if __name__=="__main__":
    sys.path.insert(0,str(Path(__file__).resolve().parents[1]));from agent_excel_logger import run_logged;run_logged("activity_proposals","activity_proposals",main)
